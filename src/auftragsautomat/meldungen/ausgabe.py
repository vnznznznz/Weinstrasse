"""Ausgabedateien je Firma und Monat: meldungen/<firma>/<jahr>-<monat>/ ...

* intrastat_eingang_idev.csv / intrastat_versendung_idev.csv  – Import ins IDEV-Formular
* meldungen_<firma>_<monat>.xlsx – Prüflisten (Intrastat, Schaumweinsteuer mit Belegliste,
  Verpackung je Materialart, Fristen, Hinweise)
* zusammenfassung.md – Checkliste zum Abarbeiten in IDEV, Zoll-Portal, Kundenportal, LUCID
"""

from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from auftragsautomat.config import Firma, Settings
from auftragsautomat.meldungen.intrastat import idev_csv_zeilen
from auftragsautomat.meldungen.modelle import (
    ARTEN_DES_GESCHAEFTS,
    BUNDESLAENDER,
    VERKEHRSZWEIGE,
    IntrastatMeldung,
    Material,
    Monatsmeldungen,
    Richtung,
)
from auftragsautomat.models import Pruefhinweis, Schweregrad

FETT = Font(bold=True)
ROT = PatternFill("solid", fgColor="F8CBAD")
GELB = PatternFill("solid", fgColor="FFE699")


def ausgabe_ordner(settings: Settings, firma: Firma, monat) -> Path:
    ordner = settings.meldungen.ausgabe_ordner / firma.id / str(monat)
    ordner.mkdir(parents=True, exist_ok=True)
    return ordner


def _d(wert: Decimal | None, stellen: int = 2) -> str:
    if wert is None:
        return ""
    return f"{wert:,.{stellen}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def schreibe_idev_csv(meldung: IntrastatMeldung, ordner: Path, settings: Settings) -> Path | None:
    if meldung.fehlanzeige:
        return None
    pfad = ordner / f"intrastat_{meldung.richtung.value}_idev.csv"
    with pfad.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\r\n")
        for zeile in idev_csv_zeilen(meldung, settings):
            w.writerow(zeile)
    return pfad


def _blatt(wb: Workbook, titel: str, kopf: list[str]):
    ws = wb.create_sheet(titel)
    ws.append(kopf)
    for z in ws[1]:
        z.font = FETT
    return ws


def _autobreite(ws) -> None:
    for spalte in ws.columns:
        breite = max((len(str(z.value)) for z in spalte if z.value is not None), default=8)
        ws.column_dimensions[get_column_letter(spalte[0].column)].width = min(max(10, breite + 2), 70)


def schreibe_xlsx(m: Monatsmeldungen, firma: Firma, ordner: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    ws = _blatt(wb, "Übersicht", ["Meldung", "Kennzahl", "Wert"])
    ie, iv, st, vp = m.intrastat_eingang, m.intrastat_versendung, m.schaumweinsteuer, m.verpackung
    zeilen = [
        ("Intrastat Eingang", "Positionen", len(ie.positionen)),
        ("Intrastat Eingang", "Rechnungsbetrag EUR", ie.summe_rechnungsbetrag),
        ("Intrastat Eingang", "Fehlanzeige", "ja" if ie.fehlanzeige else "nein"),
        ("Intrastat Versendung", "Positionen", len(iv.positionen)),
        ("Intrastat Versendung", "Rechnungsbetrag EUR", iv.summe_rechnungsbetrag),
        ("Intrastat Versendung", "Fehlanzeige", "ja" if iv.fehlanzeige else "nein"),
        ("Schaumweinsteuer", "Liter steuerpflichtig", _d(st.liter_gesamt, 3)),
        ("Schaumweinsteuer", "Steuer EUR", _d(st.steuer_gesamt_eur)),
        ("Schaumweinsteuer", "Abgabe bis", st.abgabe_bis.isoformat() if st.abgabe_bis else ""),
        ("Schaumweinsteuer", "Fällig am", st.faellig_am.isoformat() if st.faellig_am else ""),
        ("Verpackung", "Flaschen Monat (Inland)", _d(vp.flaschen_monat, 0)),
        ("Verpackung", "Flaschen Jahr bis Monat", _d(vp.flaschen_jahr_bis_monat, 0)),
    ]
    for mat in Material:
        zeilen.append(("Verpackung", f"{mat.value} kg Monat", _d(vp.kg(mat), 3)))
    for z in zeilen:
        ws.append(list(z))
    _autobreite(ws)

    for meldung in (ie, iv):
        titel = "Intrastat Eingang" if meldung.richtung == Richtung.EINGANG else "Intrastat Versendung"
        ws = _blatt(
            wb,
            titel,
            [
                "Warennummer",
                "Versendungs-/Bestimmungsmitgliedstaat",
                "Ursprungsland",
                "Bundesland",
                "Art des Geschäfts",
                "Verkehrszweig",
                "USt-IdNr. Partner",
                "Eigenmasse kg",
                "Bes. Maßeinheit (Liter)",
                "Rechnungsbetrag EUR",
                "Statistischer Wert EUR",
                "Flaschen",
                "Liter",
                "Belege",
                "Warenbezeichnung",
            ],
        )
        for p in meldung.positionen:
            ws.append(
                [
                    p.warennummer,
                    p.partner_mitgliedstaat,
                    p.ursprungsland,
                    f"{p.bundesland} {BUNDESLAENDER.get(p.bundesland, '')}".strip(),
                    f"{p.art_des_geschaefts} {ARTEN_DES_GESCHAEFTS.get(p.art_des_geschaefts, '')}",
                    f"{p.verkehrszweig} {VERKEHRSZWEIGE.get(p.verkehrszweig, '')}",
                    p.partner_ustid or "",
                    p.eigenmasse_kg,
                    p.besondere_masseinheit if p.besondere_masseinheit is not None else "",
                    p.rechnungsbetrag_eur,
                    p.statistischer_wert_eur if p.statistischer_wert_eur is not None else "",
                    float(p.flaschen),
                    float(p.liter),
                    ", ".join(p.belegnummern),
                    p.warenbezeichnung or "",
                ]
            )
        if meldung.fehlanzeige:
            ws.append(["FEHLANZEIGE – keine Warenverkehre in dieser Richtung; Fehlanzeige in IDEV abgeben."])
        _autobreite(ws)

    ws = _blatt(
        wb,
        "Schaumweinsteuer 2401",
        ["Steuerklasse", "Steuersatz EUR/hl", "Flaschen", "Liter", "Hektoliter", "Steuer EUR", "Belege"],
    )
    for z in st.zeilen:
        ws.append(
            [
                z.klasse.value,
                float(z.steuersatz_eur_je_hl),
                float(z.flaschen),
                float(z.liter),
                float(z.liter / 100),
                float(z.steuer_eur),
                ", ".join(z.belegnummern),
            ]
        )
    ws.append(
        [
            "Summe",
            "",
            float(sum((z.flaschen for z in st.zeilen), Decimal(0))),
            float(st.liter_gesamt),
            float(st.liter_gesamt / 100),
            float(st.steuer_gesamt_eur),
            "",
        ]
    )
    ws[ws.max_row][0].font = FETT
    _autobreite(ws)

    ws = _blatt(
        wb,
        "Schaumweinsteuer Belege",
        [
            "Beleg",
            "Datum",
            "Partner",
            "Land",
            "Artikel",
            "Bezeichnung",
            "Steuerklasse",
            "Flaschen",
            "Liter",
            "% vol",
            "Steuer EUR",
        ],
    )
    for b in st.belege:
        ws.append(
            [
                b.belegnummer,
                b.datum.isoformat(),
                b.partnername or "",
                b.partner_land,
                b.artikelnummer,
                b.bezeichnung or "",
                b.klasse.value if b.klasse else "",
                float(b.flaschen),
                float(b.liter),
                float(b.alkohol_vol) if b.alkohol_vol is not None else "",
                float(b.steuer_eur),
            ]
        )
    _autobreite(ws)

    ws = _blatt(
        wb,
        "Verpackung",
        [
            "Materialart (LUCID)",
            f"kg {m.monat}",
            f"kg {m.monat.jahr} bis {m.monat}",
            f"Hochrechnung kg {m.monat.jahr} (Planmenge)",
        ],
    )
    hoch = {h.material: h.kg for h in vp.hochrechnung_jahr}
    for mat in Material:
        ws.append(
            [mat.value, float(vp.kg(mat)), float(vp.kg(mat, jahr=True)), float(hoch.get(mat, Decimal(0)))]
        )
    ws.append([])
    ws.append(["Systembetreiber", vp.systembetreiber or "(fehlt)"])
    ws.append(["LUCID-Registrierungsnummer", vp.lucid_registrierungsnummer or "(fehlt)"])
    ws.append(["Flaschen Inland Monat", float(vp.flaschen_monat)])
    ws.append(["Flaschen Inland Jahr bis Monat", float(vp.flaschen_jahr_bis_monat)])
    _autobreite(ws)

    ws = _blatt(wb, "Fristen", ["Datum", "Meldung", "Bezeichnung", "Empfänger", "Hinweis"])
    for f in m.fristen:
        ws.append([f.datum.isoformat(), f.meldung, f.bezeichnung, f.empfaenger, f.hinweis or ""])
    _autobreite(ws)

    ws = _blatt(wb, "Hinweise", ["Schweregrad", "Code", "Text", "Details"])
    for h in sorted(m.alle_hinweise, key=_rang):
        ws.append(
            [h.schweregrad.value, h.code.value, h.text, "; ".join(f"{k}={v}" for k, v in h.details.items())]
        )
        if h.schweregrad == Schweregrad.FEHLER:
            for z in ws[ws.max_row]:
                z.fill = ROT
        elif h.schweregrad == Schweregrad.WARNUNG:
            for z in ws[ws.max_row]:
                z.fill = GELB
    _autobreite(ws)

    pfad = ordner / f"meldungen_{firma.id}_{m.monat}.xlsx"
    wb.save(pfad)
    return pfad


def _rang(h: Pruefhinweis) -> int:
    return {Schweregrad.FEHLER: 0, Schweregrad.WARNUNG: 1, Schweregrad.INFO: 2}[h.schweregrad]


def zusammenfassung(m: Monatsmeldungen, firma: Firma, dateien: dict[str, Path | None]) -> str:
    ie, iv, st, vp = m.intrastat_eingang, m.intrastat_versendung, m.schaumweinsteuer, m.verpackung
    fehler = [h for h in m.alle_hinweise if h.schweregrad == Schweregrad.FEHLER]
    warnungen = [h for h in m.alle_hinweise if h.schweregrad == Schweregrad.WARNUNG]
    z: list[str] = [f"# Meldungen {firma.name} – {m.monat}", ""]
    if fehler:
        z += [f"**{len(fehler)} Fehler – vor der Abgabe beheben:**", ""]
        z += [f"- [ ] {h.text}" for h in fehler] + [""]
    z += ["## 1. Intrastat (IDEV, Statistisches Bundesamt)", ""]
    z.append(
        f"- Frist: **{next((f.datum.isoformat() for f in m.fristen if f.meldung == 'intrastat'), '?')}** (10. Arbeitstag)"
    )
    for meldung, name in ((ie, "Eingang"), (iv, "Versendung")):
        if meldung.meldepflichtig is False:
            z.append(f"- {name}: nicht meldepflichtig (unter der Schwelle) – keine Abgabe.")
        elif meldung.fehlanzeige:
            z.append(f"- [ ] {name}: **Fehlanzeige** abgeben (keine Warenverkehre).")
        else:
            datei = dateien.get(f"idev_{meldung.richtung.value}")
            z.append(
                f"- [ ] {name}: {len(meldung.positionen)} Position(en), Rechnungsbetrag {meldung.summe_rechnungsbetrag} EUR. "
                f"CSV importieren: `{datei.name if datei else '-'}` oder Werte aus dem Blatt 'Intrastat {name}' eingeben."
            )
    z += ["", "## 2. Schaumweinsteuer (Zoll-Portal, Formular 2401 monatlich)", ""]
    if st.zeilen:
        z.append(
            f"- [ ] Steueranmeldung bis **{st.abgabe_bis}** abgeben; Zahlung **{_d(st.steuer_gesamt_eur)} EUR** bis **{st.faellig_am}**."
        )
        for zeile in st.zeilen:
            z.append(
                f"  - {zeile.klasse.value}: {_d(zeile.liter, 2)} l = {_d(zeile.liter / 100, 4)} hl × {_d(zeile.steuersatz_eur_je_hl)} EUR/hl = {_d(zeile.steuer_eur)} EUR"
            )
    else:
        z.append(
            "- Keine steuerpflichtigen Mengen im Monat. Bei Dauererlaubnis trotzdem Null-Anmeldung abgeben (mit dem Hauptzollamt abgestimmt?)."
        )
    z += ["", f"## 3. Verpackung ({vp.systembetreiber or 'Systembetreiber'} + LUCID)", ""]
    z.append(
        f"- Flaschen Inland im Monat: {_d(vp.flaschen_monat, 0)}; Jahr bis {m.monat}: {_d(vp.flaschen_jahr_bis_monat, 0)}"
    )
    for mat in Material:
        kg_m, kg_j = vp.kg(mat), vp.kg(mat, jahr=True)
        if kg_m or kg_j:
            z.append(f"  - {mat.value}: {_d(kg_m, 3)} kg (Monat), {_d(kg_j, 3)} kg (Jahr bis Monat)")
    turnus = firma.meldungen.verpackung.meldeturnus
    z.append(
        f"- Meldeturnus laut Vertrag: {turnus}. Jede Meldung beim Systembetreiber **identisch** in LUCID spiegeln."
    )
    z += ["", "## Fristen", ""]
    z += [f"- {f.datum.isoformat()}: {f.bezeichnung} → {f.empfaenger}" for f in m.fristen]
    if warnungen:
        z += ["", "## Warnungen", ""] + [f"- {h.text}" for h in warnungen]
    z += ["", "## Dateien", ""] + [f"- {k}: `{v}`" for k, v in dateien.items() if v]
    return "\n".join(z) + "\n"
