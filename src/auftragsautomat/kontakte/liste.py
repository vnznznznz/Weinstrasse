"""Kontaktliste "Ansprechpartner Abwicklung.xlsx" lesen (docs/AUSBAUSTUFE_2.md, Abschnitt 2.3).

Aufbau der Datei (Abschnitt 1 und 2.3 des Designs):

* Blatt "Übersicht" wird ignoriert.
* Blatt "Lieferanten": Zeile 1 Titel, dann Blöcke je Land: Landname in Spalte A, darunter die
  Kopfzeile ``Firma | Name | Position | E-Mail | Telefon``, dann die Kontaktzeilen; eine Leerzeile
  trennt die Blöcke.
* Blatt "Speditionen": Titel, Erklärung, dann je Spedition eine Tabelle mit Kopf
  ``Name | Bereich / Position | E-Mail | Telefon`` (Tabellenobjekte ``tbl_Spedition_<Name>``).
  Der Speditionsname steht **nicht** als Zelle über der Tabelle; er kommt aus dem Tabellennamen
  (lokale Datei: openpyxl ``ws.tables``; Backend: optionale Methode ``tabellen``), sonst aus der
  E-Mail-Domain bzw. dem Positionstext, abgeglichen mit den bekannten Speditionen (Spalte
  "Spedition" des Kundenblatts und ``settings.kontakte.standard_spedition``).
* Blatt "Kunden": Titel, dann Blöcke je Kunde: Kundenname in Spalte A, Kopf
  ``Produkt | Name | Position | E-Mail | Telefon | Spedition | Abwicklungsverlauf``.

Die Kopfzeilen werden über Schlüsselwörter erkannt (Groß-/Kleinschreibung, Mehrfach-Leerzeichen
und fehlende Spalten sind unerheblich); Telefonnummern dürfen als Zahl gespeichert sein. Jeder
:class:`~auftragsautomat.models.Kontakt` erhält ``quelle = "<Blatt>!<Excel-Zeile>"``, damit
``aenderungen.uebernehme`` die Zelle wiederfindet.

Lokale Dateien liest openpyxl ausschließlich **lesend** (``read_only=True, data_only=True``) auf
einer Kopie im Temp-Ordner; die Originaldatei wird nie gespeichert (Roundtrip würde Tabellen,
Kontrollkästchen und dynamische Formeln beschädigen, siehe ``docs/recherche/excel_lokal_und_vorlagen.md``).
Über ein :class:`~auftragsautomat.m365.basis.ExcelBackend` (Graph, COM, Fake) wird der benutzte
Bereich jedes Blatts gelesen (``lese_benutzten_bereich``).
"""

from __future__ import annotations

import logging
import re
import shutil
import tempfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field

from auftragsautomat.benennung import normalisiere
from auftragsautomat.config import DateiRefKonfig, Settings
from auftragsautomat.m365.basis import ExcelBackend, M365Fehler, parse_adresse
from auftragsautomat.matching.aliase import aehnlichkeit, normalisiere_name
from auftragsautomat.models import Kontakt, Kontaktgruppe

__all__ = [
    "FELDER",
    "KontakteFehler",
    "Kontaktliste",
    "bereinige_speditionsname",
    "erkenne_kopf",
    "kontaktliste_aus_settings",
    "lade_kontaktliste",
    "lies_xlsx_lokal",
    "quelle_text",
    "telefon_text",
    "zerlege_quelle",
]

logger = logging.getLogger(__name__)

#: Spaltenfelder, die in den Kopfzeilen erkannt werden.
Feld = Literal["firma", "name", "position", "email", "telefon", "produkt", "spedition", "verlauf", "land"]
FELDER: tuple[str, ...] = ("firma", "name", "position", "email", "telefon", "produkt", "spedition", "verlauf", "land")

#: Reihenfolge ist wichtig: "E-Mail" vor "Name", "Firmenname" -> firma, "Bereich / Position" -> position.
_KOPF_MUSTER: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("email", re.compile(r"e-?mail|e-?post|mail\b")),
    ("telefon", re.compile(r"telefon|tel\b|tel\.|phone|\bfon\b|durchwahl|mobil|handy")),
    ("verlauf", re.compile(r"verlauf|abwicklung|bemerkung|notiz|hinweis|kommentar")),
    ("spedition", re.compile(r"spedition|spediteur|logistik|carrier|transport")),
    ("firma", re.compile(r"firm(?:a|en)|unternehmen|company|weingut|betrieb|^kunde|^lieferant")),
    ("produkt", re.compile(r"produkt|artikel|sortiment|marke|wein\b|weine\b")),
    ("position", re.compile(r"position|bereich|funktion|rolle|abteilung|zust")),
    ("name", re.compile(r"\bname\b|ansprechpartner|kontakt|person")),
    ("land", re.compile(r"^land$|country")),
)
_MAX_KOPFZELLE = 40
_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_MEHRFACH_LEER = re.compile(r"[ \t\r\f\v\xa0]+")
_KLAMMER_ZUSATZ = re.compile(r"\s*\((?:f[üu]r|for|per|pour)\b[^)]*\)\s*", re.IGNORECASE)
_KLAMMER_ALLGEMEIN = re.compile(r"\s*\([^)]*\)\s*")
_TABELLEN_PRAEFIX = re.compile(r"^(?:tbl|tab|tabelle|table)?[_\s-]*spedition(?:en)?[_\s-]*", re.IGNORECASE)
_LAENDER_DEUTSCH = {
    "deutschland",
    "germany",
    "de",
    "d",
    "oesterreich",
    "austria",
    "at",
    "a",
    "schweiz",
    "switzerland",
    "suisse",
    "svizzera",
    "ch",
}
#: Domain-Bestandteile, aus denen sich kein Speditionsname ableiten lässt.
_GENERISCHE_LABELS = {
    "mail",
    "email",
    "e-mail",
    "post",
    "www",
    "web",
    "smtp",
    "office",
    "info",
    "example",
    "co",
    "com",
    "de",
    "eu",
    "net",
    "org",
    "it",
    "fr",
    "es",
    "at",
    "ch",
    "uk",
    "gmail",
    "googlemail",
    "outlook",
    "hotmail",
    "live",
    "gmx",
    "t-online",
    "yahoo",
    "icloud",
    "freenet",
}


class KontakteFehler(Exception):
    """Kontaktliste nicht lesbar oder Änderung nicht übernehmbar (Konfiguration/Datei kaputt)."""


# ---------------------------------------------------------------------------
# Zelltexte und Hilfen
# ---------------------------------------------------------------------------


def zelltext(wert: Any) -> str:
    """Zellinhalt als bereinigter Text: ``None``/leer -> ``""``, Zahlen ohne ``.0``, Datumswerte
    deutsch, Mehrfach-Leerzeichen zusammengefasst, Zeilenumbrüche -> ``"; "``."""
    if wert is None:
        return ""
    if isinstance(wert, bool):
        return "WAHR" if wert else "FALSCH"
    if isinstance(wert, int):
        return str(wert)
    if isinstance(wert, float):
        return str(int(wert)) if wert.is_integer() else str(wert)
    if isinstance(wert, Decimal):
        return str(int(wert)) if wert == wert.to_integral_value() else str(wert)
    if isinstance(wert, datetime):
        return wert.strftime("%d.%m.%Y %H:%M")
    if isinstance(wert, date):
        return wert.strftime("%d.%m.%Y")
    text = str(wert).replace("\r\n", "\n").replace("\r", "\n")
    zeilen = [_MEHRFACH_LEER.sub(" ", z).strip() for z in text.split("\n")]
    return "; ".join(z for z in zeilen if z)


def telefon_text(wert: Any) -> str | None:
    """Telefonnummer als Text; Zahlen (Excel speichert "+49…" gern als Zahl) werden zu Ziffernfolgen.

    Ein verlorenes ``+`` wird nicht geraten; die Nummer bleibt, wie sie in der Zelle steht.
    """
    text = zelltext(wert)
    return text or None


def _email_text(wert: Any) -> tuple[str | None, str | None]:
    """Erste gültige E-Mail-Adresse (kleingeschrieben) und ggf. ein Hinweis (ungültig/mehrere)."""
    text = zelltext(wert)
    if not text:
        return None, None
    treffer = _EMAIL_RE.findall(text)
    if not treffer:
        return None, f"Kein gültiges E-Mail-Format: '{text}'"
    hinweis = None
    if len(treffer) > 1:
        hinweis = f"Mehrere E-Mail-Adressen in einer Zelle, nur die erste wird genutzt: '{text}'"
    return treffer[0].lower(), hinweis


def bereinige_speditionsname(text: str | None) -> str:
    """"Bouché (für Rheinberg)" -> "Bouché"; Klammerzusätze und Mehrfach-Leerzeichen entfernt."""
    if not text:
        return ""
    t = _KLAMMER_ZUSATZ.sub(" ", zelltext(text))
    t = _KLAMMER_ALLGEMEIN.sub(" ", t)
    return _MEHRFACH_LEER.sub(" ", t).strip(" -–,;")


def sprache_fuer_land(land: str | None, *, standard: Literal["de", "en"] = "en") -> Literal["de", "en"]:
    """Deutsch für DE/AT/CH, sonst Englisch; ohne Land ``standard``."""
    if not land:
        return standard
    return "de" if normalisiere(land) in _LAENDER_DEUTSCH else "en"


def quelle_text(blatt: str, zeile: int) -> str:
    """Quellangabe eines Kontakts: ``"<Blatt>!<Excel-Zeile>"``."""
    return f"{blatt}!{zeile}"


def zerlege_quelle(quelle: str | None) -> tuple[str, int]:
    """``"Lieferanten!7"`` bzw. ``"'Blatt mit Leerzeichen'!7"`` -> ``("Lieferanten", 7)``."""
    if not quelle or "!" not in quelle:
        raise KontakteFehler(f"Kontakt ohne verwertbare Quelle ('{quelle}'); erwartet '<Blatt>!<Zeile>'.")
    blatt, _, zeile = quelle.rpartition("!")
    blatt = blatt.strip()
    if blatt.startswith("'") and blatt.endswith("'"):
        blatt = blatt[1:-1]
    try:
        nummer = int(zeile.strip())
    except ValueError as exc:
        raise KontakteFehler(f"Quelle '{quelle}' enthält keine Zeilennummer.") from exc
    if not blatt or nummer < 1:
        raise KontakteFehler(f"Quelle '{quelle}' ist unvollständig.")
    return blatt, nummer


def feld_aus_kopf(text: str) -> str | None:
    """Ordnet einen Kopfzellentext einem Feld zu (``"E-Mail"`` -> ``email``), sonst ``None``."""
    t = normalisiere(text)
    if not t or len(t) > _MAX_KOPFZELLE or "@" in t:
        return None
    for feld, muster in _KOPF_MUSTER:
        if muster.search(t):
            return feld
    return None


def erkenne_kopf(zellen: Sequence[str]) -> dict[str, int] | None:
    """Erkennt eine Kopfzeile: mindestens zwei bekannte Felder, darunter "Name" oder "E-Mail".
    Ergebnis: Feld -> Spaltenindex (0-basiert innerhalb der Matrix)."""
    felder: dict[str, int] = {}
    for i, text in enumerate(zellen):
        if not text:
            continue
        feld = feld_aus_kopf(text)
        if feld and feld not in felder:
            felder[feld] = i
    if len(felder) >= 2 and ("name" in felder or "email" in felder):
        return felder
    return None


# ---------------------------------------------------------------------------
# Blattstruktur
# ---------------------------------------------------------------------------


@dataclass
class _Blatt:
    name: str
    matrix: list[list[Any]]
    zeile_offset: int = 1  # Excel-Zeile der ersten Matrixzeile
    spalte_offset: int = 1  # Excel-Spalte (1 = A) der ersten Matrixspalte
    tabellen: dict[str, str] = field(default_factory=dict)  # Tabellenname -> Bereich


@dataclass
class _Block:
    label: str | None
    kopf: dict[str, int]
    kopfzeile: int  # Excel-Zeile
    zeilen: list[tuple[int, list[str]]] = field(default_factory=list)  # (Excel-Zeile, Zelltexte)


def _zeilentexte(zeile: Iterable[Any]) -> list[str]:
    return [zelltext(z) for z in zeile]


def _ist_label(zellen: Sequence[str]) -> str | None:
    """Eine Zeile mit genau einer gefüllten Zelle (Land, Kundenname) -> ihr Text."""
    gefuellt = [z for z in zellen if z]
    if len(gefuellt) != 1 or "@" in gefuellt[0] or len(gefuellt[0]) > 80:
        return None
    return gefuellt[0]


def _naechste_gefuellte(matrix: list[list[Any]], ab: int) -> list[str] | None:
    for zeile in matrix[ab:]:
        texte = _zeilentexte(zeile)
        if any(texte):
            return texte
    return None


def bloecke(blatt: _Blatt) -> list[_Block]:
    """Zerlegt ein Blatt in Blöcke (Label-Zeile, Kopfzeile, Datenzeilen). Leerzeilen beenden
    einen Block; eine einzelne gefüllte Zelle direkt vor einer Kopfzeile ist das Label."""
    ergebnis: list[_Block] = []
    label: str | None = None
    aktuell: _Block | None = None
    for i, roh in enumerate(blatt.matrix):
        zellen = _zeilentexte(roh)
        excel_zeile = blatt.zeile_offset + i
        if not any(zellen):
            aktuell = None
            label = None
            continue
        kopf = erkenne_kopf(zellen)
        if kopf is not None:
            aktuell = _Block(label=label, kopf=kopf, kopfzeile=excel_zeile)
            ergebnis.append(aktuell)
            label = None
            continue
        if aktuell is None:
            label = _ist_label(zellen)
            continue
        kandidat = _ist_label(zellen)
        if kandidat is not None:
            naechste = _naechste_gefuellte(blatt.matrix, i + 1)
            if naechste is not None and erkenne_kopf(naechste) is not None:
                # Label des nächsten Blocks ohne trennende Leerzeile
                aktuell = None
                label = kandidat
                continue
        aktuell.zeilen.append((excel_zeile, zellen))
    return ergebnis


def _wert(zellen: Sequence[str], kopf: dict[str, int], feld: str) -> str:
    index = kopf.get(feld)
    if index is None or index >= len(zellen):
        return ""
    return zellen[index]


# ---------------------------------------------------------------------------
# Ergebnismodell
# ---------------------------------------------------------------------------


class Kontaktliste(BaseModel):
    """Alle Kontakte der Liste, zusätzlich gruppiert nach Lieferant, Spedition und Kunde.

    Die Gruppen-Dicts sind nach Firmenname (wie in der Liste) geschlüsselt und enthalten dieselben
    :class:`Kontakt`-Objekte wie ``kontakte``. ``abwicklungsverlauf`` hält den Text der Spalte
    "Abwicklungsverlauf" je Kunde (das Kontakt-Modell hat dafür kein Feld).
    """

    kontakte: list[Kontakt] = Field(default_factory=list)
    lieferanten: dict[str, list[Kontakt]] = Field(default_factory=dict)
    speditionen: dict[str, list[Kontakt]] = Field(default_factory=dict)
    kunden: dict[str, list[Kontakt]] = Field(default_factory=dict)
    abwicklungsverlauf: dict[str, str] = Field(default_factory=dict)
    blaetter: dict[str, str] = Field(default_factory=dict)  # Gruppe -> tatsächlicher Blattname
    quelle: str | None = None
    hinweise: list[str] = Field(default_factory=list)

    def je_gruppe(self, gruppe: Kontaktgruppe) -> list[Kontakt]:
        return [k for k in self.kontakte if k.gruppe == gruppe]

    def kontakt_fuer_email(self, email: str | None) -> Kontakt | None:
        """Kontakt mit genau dieser E-Mail-Adresse (Groß-/Kleinschreibung egal)."""
        if not email:
            return None
        gesucht = email.strip().lower()
        for kontakt in self.kontakte:
            if kontakt.email and kontakt.email.lower() == gesucht:
                return kontakt
        return None

    def finde_spedition(self, name: str | None) -> str | None:
        """Schlüssel in ``speditionen`` für einen Speditionsnamen (normalisiert, Teilstring,
        sonst unscharf ab 85). "Hillebrand Gori" trifft "Hillebrand"."""
        return finde_schluessel(self.speditionen, bereinige_speditionsname(name))

    def kontakte_der_spedition(self, name: str | None) -> list[Kontakt]:
        schluessel = self.finde_spedition(name)
        return list(self.speditionen[schluessel]) if schluessel else []

    def bekannte_speditionen(self) -> list[str]:
        """Speditionsnamen aus dem Speditionsblatt und der Kundenspalte "Spedition" (ohne Dubletten)."""
        namen: list[str] = list(self.speditionen)
        for kontakt in self.kontakte:
            s = bereinige_speditionsname(kontakt.spedition)
            if s and finde_schluessel({n: None for n in namen}, s) is None:
                namen.append(s)
        return namen


def finde_schluessel(kandidaten: dict[str, Any], name: str | None, *, schwelle: float = 85.0) -> str | None:
    """Bester Schlüssel für einen Namen: exakt normalisiert, dann Teilstring (ab 4 Zeichen,
    ohne Leerzeichen verglichen), dann unscharf (``aehnlichkeit`` >= ``schwelle``)."""
    gesucht = normalisiere_name(name)
    if not gesucht or not kandidaten:
        return None
    normiert = {schluessel: normalisiere_name(schluessel) for schluessel in kandidaten}
    for schluessel, norm in normiert.items():
        if norm == gesucht:
            return schluessel
    kompakt = gesucht.replace(" ", "")
    if len(kompakt) >= 4:
        for schluessel, norm in normiert.items():
            kn = norm.replace(" ", "")
            if len(kn) >= 4 and (kompakt in kn or kn in kompakt):
                return schluessel
    bester, score = None, 0.0
    for schluessel, norm in normiert.items():
        s = aehnlichkeit(gesucht, norm)
        if s > score:
            bester, score = schluessel, s
    return bester if score >= schwelle else None


# ---------------------------------------------------------------------------
# Blätter lesen (lokal / Backend)
# ---------------------------------------------------------------------------


def lies_xlsx_lokal(pfad: Path) -> tuple[dict[str, list[list[Any]]], dict[str, dict[str, str]]]:
    """Liest alle Blätter (Werte) und die Tabellennamen je Blatt aus einer ``.xlsx``.

    Arbeitet auf einer Kopie im Temp-Ordner, nur lesend. Werte: ``read_only=True, data_only=True``;
    Tabellennamen (``ws.tables``) erfordern ein zweites Öffnen ohne ``read_only`` – schlägt das fehl,
    werden die Speditionsnamen aus Domain/Positionstext abgeleitet.
    """
    from openpyxl import load_workbook

    pfad = Path(pfad)
    if not pfad.is_file():
        raise KontakteFehler(f"Kontaktliste nicht gefunden: {pfad}")
    blaetter: dict[str, list[list[Any]]] = {}
    tabellen: dict[str, dict[str, str]] = {}
    with tempfile.TemporaryDirectory(prefix="auftragsautomat_kontakte_") as tmp:
        kopie = Path(tmp) / pfad.name
        shutil.copy2(pfad, kopie)
        try:
            mappe = load_workbook(kopie, read_only=True, data_only=True)
        except Exception as exc:
            raise KontakteFehler(f"Kontaktliste {pfad} nicht lesbar: {exc}") from exc
        try:
            for ws in mappe.worksheets:
                blaetter[ws.title] = [list(zeile) for zeile in ws.iter_rows(values_only=True)]
        finally:
            mappe.close()
        try:
            voll = load_workbook(kopie, read_only=False, data_only=True)
            try:
                for ws in voll.worksheets:
                    tabellen[ws.title] = {name: ref for name, ref in ws.tables.items()}
            finally:
                voll.close()
        except Exception as exc:
            logger.warning("Tabellennamen aus %s nicht lesbar (%s); Speditionsnamen werden abgeleitet.", pfad, exc)
    return blaetter, tabellen


def _blaetter_lokal(pfad: Path) -> dict[str, _Blatt]:
    werte, tabellen = lies_xlsx_lokal(pfad)
    return {name: _Blatt(name, matrix, 1, 1, tabellen.get(name, {})) for name, matrix in werte.items()}


def _blaetter_backend(backend: ExcelBackend, ref: DateiRefKonfig) -> dict[str, _Blatt]:
    try:
        namen = backend.blattnamen(ref)
    except M365Fehler as exc:
        raise KontakteFehler(f"Kontaktliste über Backend nicht lesbar: {exc}") from exc
    ergebnis: dict[str, _Blatt] = {}
    tabellen_funktion = getattr(backend, "tabellen", None)
    for name in namen:
        if _gruppe_fuer_blatt(name) is None:
            continue
        try:
            adresse, matrix = backend.lese_benutzten_bereich(ref, name)
        except M365Fehler as exc:
            raise KontakteFehler(f"Blatt '{name}' der Kontaktliste nicht lesbar: {exc}") from exc
        # TODO-PRUEFEN: Graph usedRange liefert bei leeren Randzeilen/-spalten eine Adresse ab z. B. B3;
        # der Versatz wird hier berücksichtigt, am echten Tenant einmal gegen die Excel-Zeilen prüfen.
        bereich = parse_adresse(adresse)
        tabellen: dict[str, str] = {}
        if callable(tabellen_funktion):
            # TODO-PRUEFEN: optionale Backend-Methode tabellen(datei, blatt) -> {Name: Bereich}
            # (Graph: GET …/worksheets/{name}/tables); kein vorhandenes Backend bietet sie heute an.
            try:
                tabellen = dict(tabellen_funktion(ref, name) or {})
            except Exception as exc:
                logger.warning("Tabellennamen des Blatts '%s' nicht lesbar (%s).", name, exc)
        ergebnis[name] = _Blatt(name, [list(z) for z in matrix], bereich.zeile_von, bereich.spalte_von, tabellen)
    return ergebnis


def _gruppe_fuer_blatt(name: str) -> Kontaktgruppe | None:
    n = normalisiere(name)
    if "uebersicht" in n or "overview" in n:
        return None
    if "lieferant" in n or "weingut" in n or "supplier" in n:
        return Kontaktgruppe.LIEFERANT
    if "spedition" in n or "spediteur" in n or "logistik" in n:
        return Kontaktgruppe.SPEDITION
    if "kunde" in n or "customer" in n:
        return Kontaktgruppe.KUNDE
    if "intern" in n:
        return Kontaktgruppe.INTERN
    return None


# ---------------------------------------------------------------------------
# Kontakte bauen
# ---------------------------------------------------------------------------


def _basisfelder(zellen: Sequence[str], kopf: dict[str, int], hinweise: list[str], quelle: str) -> dict[str, Any]:
    email, hinweis = _email_text(_wert(zellen, kopf, "email"))
    if hinweis:
        hinweise.append(f"{quelle}: {hinweis}")
    return {
        "name": _wert(zellen, kopf, "name") or None,
        "position": _wert(zellen, kopf, "position") or None,
        "email": email,
        "telefon": telefon_text(_wert(zellen, kopf, "telefon")),
        "quelle": quelle,
    }


def _hat_inhalt(felder: dict[str, Any]) -> bool:
    return any(felder.get(f) for f in ("name", "position", "email", "telefon"))


def _lieferanten(blatt: _Blatt, liste: Kontaktliste) -> None:
    for block in bloecke(blatt):
        land = block.label
        if land is None:
            liste.hinweise.append(f"{blatt.name}!{block.kopfzeile}: Block ohne Landangabe darüber.")
        firma_vorher = ""
        for zeile, zellen in block.zeilen:
            quelle = quelle_text(blatt.name, zeile)
            firma = _wert(zellen, block.kopf, "firma") or firma_vorher
            felder = _basisfelder(zellen, block.kopf, liste.hinweise, quelle)
            if not firma:
                liste.hinweise.append(f"{quelle}: Zeile ohne Firma übersprungen.")
                continue
            firma_vorher = firma
            if not _hat_inhalt(felder):
                liste.hinweise.append(f"{quelle}: Zeile ohne Name/Position/E-Mail/Telefon übersprungen.")
                continue
            kontakt = Kontakt(
                gruppe=Kontaktgruppe.LIEFERANT,
                firma=firma,
                land=land,
                sprache=sprache_fuer_land(land, standard="en"),
                **felder,
            )
            liste.kontakte.append(kontakt)
            liste.lieferanten.setdefault(firma, []).append(kontakt)


def _kunden(blatt: _Blatt, liste: Kontaktliste) -> None:
    for block in bloecke(blatt):
        kunde = block.label
        if kunde is None:
            kunde = f"Unbekannter Kunde ({blatt.name}!{block.kopfzeile})"
            liste.hinweise.append(f"{blatt.name}!{block.kopfzeile}: Kundenblock ohne Kundenname darüber.")
        produkt_vorher = ""
        verlauf: list[str] = []
        for zeile, zellen in block.zeilen:
            quelle = quelle_text(blatt.name, zeile)
            felder = _basisfelder(zellen, block.kopf, liste.hinweise, quelle)
            produkt = _wert(zellen, block.kopf, "produkt") or produkt_vorher or "Allgemein"
            produkt_vorher = produkt
            text_verlauf = _wert(zellen, block.kopf, "verlauf")
            if text_verlauf:
                verlauf.append(text_verlauf)
            spedition = _wert(zellen, block.kopf, "spedition") or None
            if not _hat_inhalt(felder):
                if not spedition and not text_verlauf:
                    liste.hinweise.append(f"{quelle}: Zeile ohne Angaben übersprungen.")
                    continue
            kontakt = Kontakt(
                gruppe=Kontaktgruppe.KUNDE,
                firma=kunde,
                produkt=produkt,
                spedition=spedition,
                sprache="de",
                **felder,
            )
            liste.kontakte.append(kontakt)
            liste.kunden.setdefault(kunde, []).append(kontakt)
        if verlauf:
            liste.abwicklungsverlauf[kunde] = "; ".join(verlauf)


def _name_aus_tabellenname(tabellenname: str) -> str:
    rest = _TABELLEN_PRAEFIX.sub("", tabellenname.strip())
    rest = rest or tabellenname
    rest = re.sub(r"^(?:tbl|tab|table|tabelle)[_\s-]*", "", rest, flags=re.IGNORECASE) or rest
    return _MEHRFACH_LEER.sub(" ", rest.replace("_", " ")).strip()


def _domain_label(email: str | None) -> str | None:
    if not email or "@" not in email:
        return None
    domain = email.rsplit("@", 1)[1].lower()
    labels = domain.split(".")
    for label in labels[:-1]:
        if label and label not in _GENERISCHE_LABELS:
            return label
    return None


def _bekannter_name(kandidat: str | None, bekannte: Sequence[str]) -> str | None:
    """Bekannter Speditionsname, der zum Kandidaten passt (normalisiert, Teilstring ab 4 Zeichen)."""
    kn = normalisiere_name(kandidat).replace(" ", "")
    if len(kn) < 4:
        return None
    for name in bekannte:
        bn = normalisiere_name(name).replace(" ", "")
        if len(bn) >= 4 and (bn == kn or bn in kn or kn in bn):
            return name
    return None


def _label_zu_name(label: str) -> str:
    return "-".join(teil[:1].upper() + teil[1:] for teil in label.split("-") if teil)


def _speditionsname(
    block: _Block,
    tabellen_kopfzeilen: dict[int, str],
    bekannte: Sequence[str],
    laufnummer: int,
    hinweise: list[str],
    blatt: str,
) -> str:
    quelle = f"{blatt}!{block.kopfzeile}"
    # 1. Tabellenobjekt tbl_Spedition_<Name>
    # TODO-PRUEFEN: Tabellennamen der echten Datei ("tbl_Spedition_<Name>", Unterstriche statt
    # Leerzeichen?) und die Positionstexte der Postfächer ("Postfach – Neue Aufträge …",
    # "Sendungsnachfragen, Status") an der echten "Ansprechpartner Abwicklung.xlsx" bestätigen.
    tabellenname = tabellen_kopfzeilen.get(block.kopfzeile)
    if tabellenname:
        name = _name_aus_tabellenname(tabellenname)
        return _bekannter_name(name, bekannte) or name
    # 2. E-Mail-Domain / Positionstext gegen bekannte Speditionen
    for _, zellen in block.zeilen:
        email, _ = _email_text(_wert(zellen, block.kopf, "email"))
        treffer = _bekannter_name(_domain_label(email), bekannte)
        if treffer:
            hinweise.append(f"{quelle}: Speditionsname '{treffer}' aus der E-Mail-Domain abgeleitet.")
            return treffer
    for _, zellen in block.zeilen:
        text = f"{_wert(zellen, block.kopf, 'name')} {_wert(zellen, block.kopf, 'position')}"
        text_norm = normalisiere_name(text)
        for name in bekannte:
            bn = normalisiere_name(name)
            if bn and f" {bn} " in f" {text_norm} ":
                hinweise.append(f"{quelle}: Speditionsname '{name}' aus dem Positionstext abgeleitet.")
                return name
    # 3. kurze Label-Zeile direkt über der Tabelle (laut Design normalerweise nicht vorhanden)
    if block.label and len(block.label) <= 40 and not re.search(r"[.;:!]", block.label):
        return _bekannter_name(block.label, bekannte) or block.label
    # 4. Domain-Bestandteil als Name
    for _, zellen in block.zeilen:
        email, _ = _email_text(_wert(zellen, block.kopf, "email"))
        label = _domain_label(email)
        if label:
            name = _label_zu_name(label)
            hinweise.append(f"{quelle}: Speditionsname '{name}' aus der Domain '{label}' geraten – bitte prüfen.")
            return name
    name = f"Spedition {laufnummer}"
    hinweise.append(f"{quelle}: Speditionsname nicht ermittelbar, Platzhalter '{name}'.")
    return name


def _tabellen_kopfzeilen(blatt: _Blatt) -> dict[int, str]:
    ergebnis: dict[int, str] = {}
    for name, ref in blatt.tabellen.items():
        try:
            ergebnis[parse_adresse(ref).zeile_von] = name
        except ValueError:
            logger.warning("Tabelle '%s' mit unlesbarem Bereich '%s' ignoriert.", name, ref)
    return ergebnis


def _speditionen(blatt: _Blatt, liste: Kontaktliste, bekannte: Sequence[str]) -> None:
    kopfzeilen = _tabellen_kopfzeilen(blatt)
    bekannt = list(bekannte)
    for nummer, block in enumerate(bloecke(blatt), start=1):
        name = _speditionsname(block, kopfzeilen, bekannt, nummer, liste.hinweise, blatt.name)
        vorhandener = finde_schluessel(liste.speditionen, name)
        if vorhandener is not None:
            name = vorhandener
        elif _bekannter_name(name, bekannt) is None:
            bekannt.append(name)
        for zeile, zellen in block.zeilen:
            quelle = quelle_text(blatt.name, zeile)
            felder = _basisfelder(zellen, block.kopf, liste.hinweise, quelle)
            if not _hat_inhalt(felder):
                liste.hinweise.append(f"{quelle}: Zeile ohne Angaben übersprungen.")
                continue
            kontakt = Kontakt(gruppe=Kontaktgruppe.SPEDITION, firma=name, sprache="de", **felder)
            liste.kontakte.append(kontakt)
            liste.speditionen.setdefault(name, []).append(kontakt)


def _intern(blatt: _Blatt, liste: Kontaktliste) -> None:
    for block in bloecke(blatt):
        firma = block.label or "Intern"
        for zeile, zellen in block.zeilen:
            quelle = quelle_text(blatt.name, zeile)
            felder = _basisfelder(zellen, block.kopf, liste.hinweise, quelle)
            if not _hat_inhalt(felder):
                continue
            liste.kontakte.append(Kontakt(gruppe=Kontaktgruppe.INTERN, firma=_wert(zellen, block.kopf, "firma") or firma, **felder))


def _baue_kontaktliste(
    blaetter: dict[str, _Blatt],
    quelle: str,
    settings: Settings | None,
    bekannte_speditionen: Sequence[str],
) -> Kontaktliste:
    liste = Kontaktliste(quelle=quelle)
    nach_gruppe: dict[Kontaktgruppe, _Blatt] = {}
    for name, blatt in blaetter.items():
        gruppe = _gruppe_fuer_blatt(name)
        if gruppe is None:
            logger.debug("Blatt '%s' der Kontaktliste wird ignoriert.", name)
            continue
        if gruppe in nach_gruppe:
            liste.hinweise.append(f"Mehrere Blätter für {gruppe.value}: '{name}' wird ignoriert.")
            continue
        nach_gruppe[gruppe] = blatt
        liste.blaetter[gruppe.value] = name

    # Kunden zuerst: ihre Spalte "Spedition" liefert bekannte Speditionsnamen.
    if Kontaktgruppe.KUNDE in nach_gruppe:
        _kunden(nach_gruppe[Kontaktgruppe.KUNDE], liste)
    else:
        liste.hinweise.append("Kein Blatt 'Kunden' gefunden.")
    if Kontaktgruppe.LIEFERANT in nach_gruppe:
        _lieferanten(nach_gruppe[Kontaktgruppe.LIEFERANT], liste)
    else:
        liste.hinweise.append("Kein Blatt 'Lieferanten' gefunden.")

    bekannte: list[str] = []
    for name in list(bekannte_speditionen) + ([settings.kontakte.standard_spedition] if settings else []):
        if name and name not in bekannte:
            bekannte.append(name)
    for name in liste.bekannte_speditionen():
        if name not in bekannte:
            bekannte.append(name)
    if Kontaktgruppe.SPEDITION in nach_gruppe:
        _speditionen(nach_gruppe[Kontaktgruppe.SPEDITION], liste, bekannte)
    else:
        liste.hinweise.append("Kein Blatt 'Speditionen' gefunden.")
    if Kontaktgruppe.INTERN in nach_gruppe:
        _intern(nach_gruppe[Kontaktgruppe.INTERN], liste)

    for hinweis in liste.hinweise:
        logger.info("Kontaktliste: %s", hinweis)
    logger.info(
        "Kontaktliste %s gelesen: %d Kontakte (%d Lieferanten, %d Speditionen, %d Kunden)",
        quelle,
        len(liste.kontakte),
        len(liste.lieferanten),
        len(liste.speditionen),
        len(liste.kunden),
    )
    return liste


# ---------------------------------------------------------------------------
# Öffentliche API
# ---------------------------------------------------------------------------


def lade_kontaktliste(
    quelle: Path | str | tuple[ExcelBackend, DateiRefKonfig],
    *,
    settings: Settings | None = None,
    bekannte_speditionen: Sequence[str] = (),
) -> Kontaktliste:
    """Liest die Kontaktliste aus einer lokalen ``.xlsx`` (Pfad) oder über ``(ExcelBackend, DateiRefKonfig)``.

    ``settings`` liefert ``kontakte.standard_spedition`` als bekannten Speditionsnamen für die
    Namensableitung; ``bekannte_speditionen`` ergänzt weitere. Fehlende Blätter, übersprungene
    Zeilen und abgeleitete Namen stehen in ``Kontaktliste.hinweise``; eine kaputte oder fehlende
    Datei ergibt :class:`KontakteFehler`.
    """
    if isinstance(quelle, tuple):
        backend, ref = quelle
        blaetter = _blaetter_backend(backend, ref)
        quelle_name = str(ref.lokaler_pfad or ref.web_url or f"{ref.drive_id}/{ref.item_id}" or "Backend")
    else:
        pfad = Path(quelle)
        blaetter = _blaetter_lokal(pfad)
        quelle_name = str(pfad)
    return _baue_kontaktliste(blaetter, quelle_name, settings, bekannte_speditionen)


def kontaktliste_aus_settings(settings: Settings, excel: ExcelBackend | None = None) -> Kontaktliste:
    """Kontaktliste laut ``settings.m365.kontaktliste_datei``: mit ``excel`` über das Backend,
    sonst über ``lokaler_pfad``. Ohne Konfiguration :class:`KontakteFehler`."""
    ref = settings.m365.kontaktliste_datei
    if excel is not None and (ref.drive_id or ref.web_url or ref.lokaler_pfad):
        return lade_kontaktliste((excel, ref), settings=settings)
    if ref.lokaler_pfad:
        return lade_kontaktliste(Path(ref.lokaler_pfad), settings=settings)
    raise KontakteFehler("settings.m365.kontaktliste_datei ist nicht konfiguriert (web_url oder lokaler_pfad).")
