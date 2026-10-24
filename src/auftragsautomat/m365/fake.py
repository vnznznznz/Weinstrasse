"""Fake-Backend für Tests und den Dry Run unter Linux (``m365.backend: fake``).

* :class:`FakeExcelBackend` – Arbeitsmappen im Speicher (``{datei_schluessel: {blatt: matrix}}``),
  optional einmalig aus einer ``.xlsx`` geladen (openpyxl ``read_only=True, data_only=True`` auf einer
  **Kopie** der Datei; es wird nie in die Datei geschrieben). ``schreibe_bereich`` wendet
  ``None = unverändert`` an, protokolliert jeden Schreibzugriff (``schreibprotokoll``) und kann die
  Änderungen als CSV-Vorschau exportieren (:meth:`FakeExcelBackend.vorschau_csv`).
* :class:`FakeMailBackend` – Entwürfe als ``.eml`` (``email.message.EmailMessage``: HTML + Klartext als
  Alternative, Anhänge, Kopfzeilen ``X-Unsent: 1`` und gültige ``Message-ID``, damit das klassische
  Outlook die Datei per Doppelklick als bearbeitbaren Entwurf öffnet) im Ordner
  ``kommunikation.entwuerfe_ordner``; Posteingang aus einer injizierten Liste von :class:`Mail`;
  Anhänge aus einem Fixture-Ordner. :func:`mail_aus_eml` baut aus einer ``.eml`` ein ``Mail``-Objekt
  für Posteingangs-Fixtures.
* :class:`FakeDateiBackend` – ein lokaler Ordner (``DateiRefKonfig.lokaler_pfad`` oder ``wurzel``).

Dateinamen der ``.eml``-Dateien sind technische Namen (``<zeit>_<typ>_<bereinigter Betreff>.eml``);
Belege, Anhänge und Ablage werden ausschließlich über ``benennung.py`` benannt.
"""

from __future__ import annotations

import csv
import email
import email.policy
import logging
import mimetypes
import shutil
import tempfile
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from email.message import EmailMessage
from email.utils import formatdate, getaddresses, make_msgid, parsedate_to_datetime
from pathlib import Path
from typing import Any

from auftragsautomat.config import DateiRefKonfig
from auftragsautomat.m365.basis import (
    DateiInfo,
    M365Fehler,
    M365NichtGefunden,
    Mail,
    MailAnhang,
    MailKopf,
    Zellbereich,
    datei_schluessel,
    html_zu_text,
    parse_adresse,
    sicherer_dateiname,
    spaltenbuchstabe,
)
from auftragsautomat.models import EmailEntwurf

__all__ = [
    "FakeDateiBackend",
    "FakeExcelBackend",
    "FakeMailBackend",
    "Schreibzugriff",
    "lade_xlsx",
    "mail_aus_eml",
]

logger = logging.getLogger(__name__)

STANDARD_SCHLUESSEL = "standard"
ENTWURF_DOMAIN = "auftragsautomat.invalid"  # reservierte TLD, nie eine echte Adresse


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def lade_xlsx(pfad: Path) -> dict[str, list[list[Any]]]:
    """Liest alle Blätter einer ``.xlsx`` als Matrizen (Werte, keine Formeln; leere Zellen ``""``).

    openpyxl arbeitet auf einer Kopie im Temp-Ordner (OneDrive-Dateien beim Lesen nicht sperren) und
    nur lesend; ein Roundtrip über openpyxl würde Kontrollkästchen und dynamische Arrays zerstören.
    """
    from openpyxl import load_workbook

    pfad = Path(pfad)
    if not pfad.is_file():
        raise M365NichtGefunden(f"Excel-Datei {pfad} nicht gefunden.")
    with tempfile.TemporaryDirectory(prefix="auftragsautomat_xlsx_") as tmp:
        kopie = Path(tmp) / pfad.name
        shutil.copy2(pfad, kopie)
        mappe = load_workbook(kopie, read_only=True, data_only=True)
        try:
            blaetter: dict[str, list[list[Any]]] = {}
            for ws in mappe.worksheets:
                zeilen = [["" if z is None else z for z in zeile] for zeile in ws.iter_rows(values_only=True)]
                blaetter[ws.title] = zeilen
            return blaetter
        finally:
            mappe.close()


@dataclass
class Schreibzugriff:
    zeitpunkt: datetime
    datei: str
    blatt: str
    adresse: str
    werte: list[list[Any]]
    zahlenformat: list[list[str | None]] | None = None
    zellen: list[tuple[str, Any, str | None]] = field(default_factory=list)  # (Zelle, Wert, Format)


class FakeExcelBackend:
    """In-Memory-Arbeitsmappen; Protokoll aller Schreibzugriffe."""

    def __init__(
        self,
        arbeitsmappen: dict[str, dict[str, list[list[Any]]]] | None = None,
        *,
        aus_xlsx: Path | None = None,
        jetzt: Callable[[], datetime] | None = None,
    ) -> None:
        self.mappen: dict[str, dict[str, list[list[Any]]]] = {}
        for schluessel, blaetter in (arbeitsmappen or {}).items():
            self.mappen[schluessel] = {name: [list(z) for z in zeilen] for name, zeilen in blaetter.items()}
        if aus_xlsx is not None:
            geladen = lade_xlsx(Path(aus_xlsx))
            self.mappen[str(Path(aus_xlsx).resolve())] = geladen
            self.mappen.setdefault(STANDARD_SCHLUESSEL, geladen)
        self.schreibprotokoll: list[Schreibzugriff] = []
        self.zahlenformate: dict[tuple[str, str, str], str] = {}  # (Datei, Blatt, Zelle) -> Format
        self._jetzt = jetzt or (lambda: datetime.now(UTC))

    # -- Zugriff auf Mappen/Blätter -----------------------------------------------------

    def _mappe(self, datei: DateiRefKonfig) -> dict[str, list[list[Any]]]:
        schluessel = datei_schluessel(datei)
        if schluessel in self.mappen:
            return self.mappen[schluessel]
        if datei.lokaler_pfad is not None:
            pfad = Path(datei.lokaler_pfad)
            if pfad.is_file():
                self.mappen[schluessel] = lade_xlsx(pfad)
                return self.mappen[schluessel]
        if STANDARD_SCHLUESSEL in self.mappen and not (datei.drive_id or datei.web_url):
            return self.mappen[STANDARD_SCHLUESSEL]
        raise M365NichtGefunden(f"Fake-Excel: keine Arbeitsmappe für '{schluessel}' hinterlegt.")

    def _blatt(self, datei: DateiRefKonfig, blatt: str) -> tuple[str, list[list[Any]]]:
        mappe = self._mappe(datei)
        name = blatt or datei.blatt or next(iter(mappe), "")
        if name not in mappe:
            raise M365NichtGefunden(f"Fake-Excel: Blatt '{name}' nicht vorhanden (vorhanden: {', '.join(mappe) or '-'}).")
        return name, mappe[name]

    def werte(self, datei: DateiRefKonfig, blatt: str) -> list[list[Any]]:
        """Die komplette Matrix eines Blatts (für Tests)."""
        return self._blatt(datei, blatt)[1]

    # -- ExcelBackend -------------------------------------------------------------

    def blattnamen(self, datei: DateiRefKonfig) -> list[str]:
        return list(self._mappe(datei))

    def lese_bereich(self, datei: DateiRefKonfig, blatt: str, adresse: str) -> list[list[Any]]:
        _, matrix = self._blatt(datei, blatt)
        bereich = parse_adresse(adresse)
        ergebnis: list[list[Any]] = []
        for zi in range(bereich.zeile_von - 1, bereich.zeile_bis):
            zeile = matrix[zi] if zi < len(matrix) else []
            ergebnis.append(
                [zeile[si] if si < len(zeile) and zeile[si] is not None else "" for si in range(bereich.spalte_von - 1, bereich.spalte_bis)]
            )
        return ergebnis

    def lese_benutzten_bereich(self, datei: DateiRefKonfig, blatt: str) -> tuple[str, list[list[Any]]]:
        name, matrix = self._blatt(datei, blatt)
        max_zeile = 0
        max_spalte = 0
        for zi, zeile in enumerate(matrix, start=1):
            for si, wert in enumerate(zeile, start=1):
                if wert not in (None, ""):
                    max_zeile = max(max_zeile, zi)
                    max_spalte = max(max_spalte, si)
        if max_zeile == 0:
            return Zellbereich(1, 1, 1, 1, name).mit_blatt(), [[""]]
        bereich = Zellbereich(1, 1, max_zeile, max_spalte, name)
        return bereich.mit_blatt(), self.lese_bereich(datei, name, bereich.adresse)

    def schreibe_bereich(
        self,
        datei: DateiRefKonfig,
        blatt: str,
        adresse: str,
        werte: list[list[Any | None]],
        *,
        zahlenformat: list[list[str | None]] | None = None,
    ) -> None:
        name, matrix = self._blatt(datei, blatt)
        bereich = parse_adresse(adresse)
        if len(werte) != bereich.zeilen or any(len(z) != bereich.spalten for z in werte):
            raise ValueError(f"werte passt nicht zum Bereich {adresse} ({bereich.zeilen}x{bereich.spalten}).")
        if zahlenformat is not None and (
            len(zahlenformat) != bereich.zeilen or any(len(z) != bereich.spalten for z in zahlenformat)
        ):
            raise ValueError(f"zahlenformat passt nicht zum Bereich {adresse}.")
        while len(matrix) < bereich.zeile_bis:
            matrix.append([])
        zugriff = Schreibzugriff(self._jetzt(), datei_schluessel(datei), name, bereich.adresse, [list(z) for z in werte], zahlenformat)
        for zi, zeile in enumerate(werte):
            ziel = matrix[bereich.zeile_von - 1 + zi]
            while len(ziel) < bereich.spalte_bis:
                ziel.append("")
            for si, wert in enumerate(zeile):
                fmt = zahlenformat[zi][si] if zahlenformat else None
                zelle = f"{spaltenbuchstabe(bereich.spalte_von + si)}{bereich.zeile_von + zi}"
                if fmt:
                    self.zahlenformate[(zugriff.datei, name, zelle)] = fmt
                if wert is None:
                    continue
                ziel[bereich.spalte_von - 1 + si] = wert
                zugriff.zellen.append((zelle, wert, fmt))
        self.schreibprotokoll.append(zugriff)
        logger.info("Fake-Excel %s!%s: %s Zellen geschrieben", name, bereich.adresse, len(zugriff.zellen))

    # -- Vorschau --------------------------------------------------------------------

    def vorschau_csv(self, pfad: Path, *, anhaengen: bool = False) -> Path:
        """Schreibt alle protokollierten Zelländerungen als CSV (UTF-8 mit BOM, ``;``) – eine Zeile je Zelle.

        Mit ``anhaengen`` werden die Zeilen an eine vorhandene Vorschau angehängt (ohne zweite Kopfzeile), damit
        mehrere Dry-Run-Buchungen nacheinander ihre Vorschau behalten.
        """
        pfad = Path(pfad)
        pfad.parent.mkdir(parents=True, exist_ok=True)
        weiter = anhaengen and pfad.is_file() and pfad.stat().st_size > 0
        with pfad.open("a" if weiter else "w", encoding="utf-8" if weiter else "utf-8-sig", newline="") as f:
            schreiber = csv.writer(f, delimiter=";")
            if not weiter:
                schreiber.writerow(["Zeitpunkt", "Datei", "Blatt", "Zelle", "Wert", "Zahlenformat"])
            for zugriff in self.schreibprotokoll:
                for zelle, wert, fmt in zugriff.zellen:
                    schreiber.writerow([zugriff.zeitpunkt.isoformat(timespec="seconds"), zugriff.datei, zugriff.blatt, zelle, _csv_wert(wert), fmt or ""])
        return pfad


def _csv_wert(wert: Any) -> str:
    if isinstance(wert, bool):
        return "WAHR" if wert else "FALSCH"
    if isinstance(wert, datetime):
        return wert.strftime("%d.%m.%Y %H:%M")
    if hasattr(wert, "strftime"):
        return wert.strftime("%d.%m.%Y")
    return str(wert)


# ---------------------------------------------------------------------------
# Mail
# ---------------------------------------------------------------------------


def _aware(zeit: datetime) -> datetime:
    return zeit.astimezone() if zeit.tzinfo is None else zeit


def mail_aus_eml(pfad: Path, *, mail_id: str | None = None, ordner: str | None = "inbox") -> Mail:
    """Baut aus einer ``.eml`` (Fixture) ein :class:`Mail`; Anhang-IDs sind die Anhangsnamen."""
    pfad = Path(pfad)
    nachricht = email.message_from_bytes(pfad.read_bytes(), policy=email.policy.default)
    von = getaddresses([str(nachricht.get("From") or "")])
    absender, absender_name = ("", None)
    if von:
        absender_name, absender = von[0][0] or None, von[0][1]
    try:
        empfangen = parsedate_to_datetime(str(nachricht.get("Date")))
    except (TypeError, ValueError):
        empfangen = datetime.fromtimestamp(pfad.stat().st_mtime, tz=UTC)
    text_teil = nachricht.get_body(preferencelist=("plain",))
    html_teil = nachricht.get_body(preferencelist=("html",))
    html = html_teil.get_content() if html_teil is not None else None
    text = text_teil.get_content() if text_teil is not None else (html_zu_text(html) if html else None)
    anhaenge = [
        MailAnhang(id=str(teil.get_filename()), name=str(teil.get_filename()), content_type=teil.get_content_type())
        for teil in nachricht.iter_attachments()
        if teil.get_filename()
    ]
    kopf = MailKopf(
        id=mail_id or str(nachricht.get("Message-ID") or pfad.stem),
        betreff=str(nachricht.get("Subject") or ""),
        absender=absender,
        absender_name=absender_name,
        empfangen=_aware(empfangen),
        hat_anhaenge=bool(anhaenge),
        ordner=ordner,
    )
    return Mail(
        kopf=kopf,
        text=text,
        html=html,
        an=[a for _, a in getaddresses([str(nachricht.get("To") or "")]) if a],
        cc=[a for _, a in getaddresses([str(nachricht.get("Cc") or "")]) if a],
        anhaenge=anhaenge,
    )


class FakeMailBackend:
    """Entwürfe als ``.eml``, Posteingang aus einer Liste. Es wird NIE gesendet."""

    def __init__(
        self,
        entwuerfe_ordner: Path,
        posteingang: list[Mail] | None = None,
        *,
        anhaenge_ordner: Path | None = None,
        jetzt: Callable[[], datetime] | None = None,
    ) -> None:
        self.entwuerfe_ordner = Path(entwuerfe_ordner)
        self.posteingang: list[Mail] = list(posteingang or [])
        self.anhaenge_ordner = Path(anhaenge_ordner) if anhaenge_ordner else None
        self.entwuerfe: list[EmailEntwurf] = []
        self._jetzt = jetzt or (lambda: datetime.now().astimezone())

    # -- Entwurf -----------------------------------------------------------------

    def erstelle_entwurf(self, entwurf: EmailEntwurf) -> str:
        jetzt = self._jetzt()
        nachricht = EmailMessage(policy=email.policy.SMTP)
        nachricht["Subject"] = entwurf.betreff
        if entwurf.absender:
            nachricht["From"] = entwurf.absender
        if entwurf.an:
            nachricht["To"] = ", ".join(entwurf.an)
        if entwurf.cc:
            nachricht["Cc"] = ", ".join(entwurf.cc)
        nachricht["Date"] = formatdate(jetzt.timestamp(), localtime=True)
        domain = entwurf.absender.rsplit("@", 1)[-1] if entwurf.absender and "@" in entwurf.absender else ENTWURF_DOMAIN
        nachricht["Message-ID"] = make_msgid(domain=domain)
        nachricht["X-Unsent"] = "1"
        nachricht["X-Auftragsautomat-Typ"] = entwurf.typ.value
        nachricht.set_content(entwurf.text_plain or html_zu_text(entwurf.text_html))
        nachricht.add_alternative(entwurf.text_html, subtype="html")
        for pfad in entwurf.anhaenge:
            pfad = Path(pfad)
            if not pfad.is_file():
                hinweis = f"Anhang fehlt und wurde nicht beigefügt: {pfad}"
                logger.warning(hinweis)
                entwurf.hinweise.append(hinweis)
                continue
            typ = mimetypes.guess_type(pfad.name)[0] or "application/octet-stream"
            haupt, _, unter = typ.partition("/")
            nachricht.add_attachment(pfad.read_bytes(), maintype=haupt, subtype=unter, filename=pfad.name)
        self.entwuerfe_ordner.mkdir(parents=True, exist_ok=True)
        basisname = f"{jetzt:%Y%m%d_%H%M%S}_{entwurf.typ.value}_{sicherer_dateiname(entwurf.betreff, maximal=80)}"
        ziel = self.entwuerfe_ordner / f"{basisname}.eml"
        laufnr = 2
        while ziel.exists():
            ziel = self.entwuerfe_ordner / f"{basisname} ({laufnr}).eml"
            laufnr += 1
        ziel.write_bytes(nachricht.as_bytes())
        entwurf.speicherort = str(ziel)
        entwurf.erstellt_am = entwurf.erstellt_am or jetzt
        self.entwuerfe.append(entwurf)
        logger.info("Entwurf als .eml gespeichert: %s", ziel)
        return str(ziel)

    # -- Posteingang --------------------------------------------------------------

    def suche_mails(
        self,
        *,
        seit: datetime,
        absender: str | None = None,
        betreff_enthaelt: str | None = None,
        ordner: str | None = None,
        maximal: int = 50,
    ) -> list[MailKopf]:
        seit = _aware(seit)
        treffer: list[MailKopf] = []
        for mail in self.posteingang:
            kopf = mail.kopf
            if _aware(kopf.empfangen) < seit:
                continue
            if absender and kopf.absender.lower() != absender.lower():
                continue
            if betreff_enthaelt and betreff_enthaelt.lower() not in kopf.betreff.lower():
                continue
            if ordner and (kopf.ordner or "inbox").lower() != ordner.lower():
                continue
            treffer.append(kopf)
        treffer.sort(key=lambda k: _aware(k.empfangen), reverse=True)
        return treffer[:maximal]

    def _mail(self, mail_id: str) -> Mail:
        for mail in self.posteingang:
            if mail.kopf.id == mail_id:
                return mail
        raise M365NichtGefunden(f"Fake-Posteingang: Nachricht {mail_id!r} nicht vorhanden.")

    def lese_mail(self, mail_id: str) -> Mail:
        return self._mail(mail_id).model_copy(deep=True)

    def speichere_anhang(self, mail_id: str, anhang_id: str, ziel: Path) -> Path:
        mail = self._mail(mail_id)
        anhang = next((a for a in mail.anhaenge if a.id == anhang_id or a.name == anhang_id), None)
        if anhang is None:
            raise M365NichtGefunden(f"Fake-Posteingang: Anhang {anhang_id!r} in {mail_id!r} nicht vorhanden.")
        kandidaten = [Path(anhang.id)]
        if self.anhaenge_ordner is not None:
            kandidaten += [self.anhaenge_ordner / anhang.name, self.anhaenge_ordner / anhang.id]
        quelle = next((k for k in kandidaten if k.is_file()), None)
        if quelle is None:
            raise M365NichtGefunden(f"Fake-Posteingang: Datei für Anhang {anhang.name!r} nicht gefunden.")
        ziel = Path(ziel)
        if ziel.is_dir():
            ziel = ziel / anhang.name
        ziel.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(quelle, ziel)
        return ziel


# ---------------------------------------------------------------------------
# Dateien
# ---------------------------------------------------------------------------


class FakeDateiBackend:
    """Ordner/Dateien in einem lokalen Ordner (``ref.lokaler_pfad`` oder ``wurzel``)."""

    def __init__(self, wurzel: Path) -> None:
        self.wurzel = Path(wurzel)

    def _basis(self, ref: DateiRefKonfig) -> Path:
        return Path(ref.lokaler_pfad) if ref.lokaler_pfad is not None else self.wurzel

    def _pfad(self, ref: DateiRefKonfig, relativ: str) -> Path:
        basis = self._basis(ref)
        ziel = (basis / relativ.strip("/")).resolve() if relativ.strip("/") else basis.resolve()
        if ziel != basis.resolve() and basis.resolve() not in ziel.parents:
            raise M365Fehler(f"Pfad '{relativ}' verlässt den Ordner {basis}.")
        return ziel

    def existiert(self, ref: DateiRefKonfig, relativ: str = "") -> bool:
        return self._pfad(ref, relativ).exists()

    def ordner_anlegen(self, ref: DateiRefKonfig, relativ: str) -> str:
        ziel = self._pfad(ref, relativ)
        if ziel.exists() and not ziel.is_dir():
            raise M365Fehler(f"'{relativ}' existiert bereits als Datei, nicht als Ordner.")
        ziel.mkdir(parents=True, exist_ok=True)
        return str(ziel)

    def liste(self, ref: DateiRefKonfig, relativ: str = "") -> list[DateiInfo]:
        ordner = self._pfad(ref, relativ)
        if not ordner.is_dir():
            raise M365NichtGefunden(f"Ordner '{relativ}' nicht gefunden unter {self._basis(ref)}.")
        eintraege: list[DateiInfo] = []
        praefix = relativ.strip("/")
        for eintrag in sorted(ordner.iterdir(), key=lambda p: p.name.lower()):
            stat = eintrag.stat()
            eintraege.append(
                DateiInfo(
                    name=eintrag.name,
                    pfad=f"{praefix}/{eintrag.name}" if praefix else eintrag.name,
                    ist_ordner=eintrag.is_dir(),
                    groesse=None if eintrag.is_dir() else stat.st_size,
                    geaendert=datetime.fromtimestamp(stat.st_mtime, tz=UTC),
                    id=str(eintrag),
                )
            )
        return eintraege

    def hochladen(self, lokal: Path, ref: DateiRefKonfig, relativ: str) -> str:
        lokal = Path(lokal)
        if not lokal.is_file():
            raise M365Fehler(f"Lokale Datei fehlt: {lokal}")
        rel = relativ.strip()
        if not rel or rel.endswith("/"):
            rel = rel + lokal.name
        ziel = self._pfad(ref, rel)
        ziel.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(lokal, ziel)
        return str(ziel)

    def herunterladen(self, ref: DateiRefKonfig, relativ: str, ziel: Path) -> Path:
        quelle = self._pfad(ref, relativ)
        if not quelle.is_file():
            raise M365NichtGefunden(f"Datei '{relativ}' nicht gefunden unter {self._basis(ref)}.")
        ziel = Path(ziel)
        if ziel.is_dir():
            ziel = ziel / quelle.name
        ziel.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(quelle, ziel)
        return ziel

    def schreibe_text(self, ref: DateiRefKonfig, relativ: str, text: str) -> str:
        ziel = self._pfad(ref, relativ)
        ziel.parent.mkdir(parents=True, exist_ok=True)
        ziel.write_text(text, encoding="utf-8")
        return str(ziel)
