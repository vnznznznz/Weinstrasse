"""Tests für m365.fake (Excel im Speicher, .eml-Entwürfe, lokaler Ordner) und m365.fabrik."""

from __future__ import annotations

import email
import email.policy
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from auftragsautomat.config import DateiRefKonfig, Settings
from auftragsautomat.m365 import fabrik
from auftragsautomat.m365.basis import M365Fehler, M365NichtGefunden, Mail, MailAnhang, MailKopf
from auftragsautomat.m365.fake import (
    FakeDateiBackend,
    FakeExcelBackend,
    FakeMailBackend,
    lade_xlsx,
    mail_aus_eml,
)
from auftragsautomat.models import EmailEntwurf, EntwurfTyp

# ---------------------------------------------------------------- Excel

KOPF = ["KW", "Datum", "Tag", "Lieferant", "Kunde", "Auftragsdatum", "Verladedatum", "Lieferant bestätigt", "Spedition bestätigt", "Abgeholt", "Bemerkung"]
BLAETTER = {
    "Fixterminüberwachung": [
        [],
        [],
        [],
        KOPF,
        [43, date(2026, 10, 19), "Mo", "Cide", "Beispiel", date(2026, 10, 1), "", True, False, False, ""],
        [43, date(2026, 10, 19), "Mo", "", "", "", "", "", "", "", ""],
    ],
    "Ressourcen für Auto DropDown": [["Cide"], ["Bollina"]],
}
REF = DateiRefKonfig(blatt="Fixterminüberwachung")


@pytest.fixture()
def excel() -> FakeExcelBackend:
    return FakeExcelBackend({"standard": BLAETTER}, jetzt=lambda: datetime(2026, 10, 18, 12, 0, tzinfo=UTC))


def test_fake_excel_lesen(excel: FakeExcelBackend):
    assert excel.blattnamen(REF) == ["Fixterminüberwachung", "Ressourcen für Auto DropDown"]
    assert excel.lese_bereich(REF, "Fixterminüberwachung", "D5:F5") == [["Cide", "Beispiel", date(2026, 10, 1)]]
    assert excel.lese_bereich(REF, "", "A1:B2") == [["", ""], ["", ""]]  # leeres Blatt -> datei.blatt
    assert excel.lese_bereich(REF, "Fixterminüberwachung", "K6:M7") == [["", "", ""], ["", "", ""]]  # außerhalb -> leer
    adresse, werte = excel.lese_benutzten_bereich(REF, "Fixterminüberwachung")
    assert adresse == "'Fixterminüberwachung'!A1:K6" and werte[3] == KOPF and werte[0] == [""] * 11
    assert excel.lese_benutzten_bereich(REF, "Ressourcen für Auto DropDown") == ("'Ressourcen für Auto DropDown'!A1:A2", [["Cide"], ["Bollina"]])
    assert BLAETTER["Fixterminüberwachung"][4][3] == "Cide"  # Konstruktor kopiert, Original unberührt
    with pytest.raises(M365NichtGefunden, match="Blatt"):
        excel.lese_bereich(REF, "Falsch", "A1")
    with pytest.raises(M365NichtGefunden, match="Arbeitsmappe"):
        excel.lese_bereich(DateiRefKonfig(drive_id="d", item_id="i"), "Fixterminüberwachung", "A1")


def test_fake_excel_schreiben_none_unveraendert_und_protokoll(excel: FakeExcelBackend, tmp_path: Path):
    excel.schreibe_bereich(
        REF, "Fixterminüberwachung", "D6:K6",
        [["Bollina", "Beispiel", date(2026, 10, 2), None, False, Decimal("1"), None, "AB 123"]],
        zahlenformat=[[None, None, "DD.MM.YYYY", None, None, None, None, None]],
    )
    zeile = excel.werte(REF, "Fixterminüberwachung")[5]
    assert zeile[3:11] == ["Bollina", "Beispiel", date(2026, 10, 2), "", False, Decimal("1"), "", "AB 123"]
    assert zeile[0:3] == [43, date(2026, 10, 19), "Mo"]  # Formelspalten A-C unberührt
    # Zelle leeren und Blatt erweitern
    excel.schreibe_bereich(REF, "Fixterminüberwachung", "D5:D5", [[""]])
    excel.schreibe_bereich(REF, "Fixterminüberwachung", "M9", [["weit draußen"]])
    assert excel.werte(REF, "Fixterminüberwachung")[4][3] == "" and excel.werte(REF, "Fixterminüberwachung")[8][12] == "weit draußen"
    assert len(excel.schreibprotokoll) == 3
    erster = excel.schreibprotokoll[0]
    assert (erster.blatt, erster.adresse, erster.datei) == ("Fixterminüberwachung", "D6:K6", "standard")
    assert [z for z, _, _ in erster.zellen] == ["D6", "E6", "F6", "H6", "I6", "K6"]  # G6 und J6 = None
    assert ("F6", date(2026, 10, 2), "DD.MM.YYYY") in erster.zellen
    assert excel.zahlenformate[("standard", "Fixterminüberwachung", "F6")] == "DD.MM.YYYY"
    csv = excel.vorschau_csv(tmp_path / "vorschau" / "fixtermine_vorschau.csv")
    inhalt = csv.read_bytes()
    assert inhalt.startswith("﻿".encode())
    zeilen = inhalt.decode("utf-8-sig").splitlines()
    assert zeilen[0] == "Zeitpunkt;Datei;Blatt;Zelle;Wert;Zahlenformat"
    assert "2026-10-18T12:00:00+00:00;standard;Fixterminüberwachung;F6;02.10.2026;DD.MM.YYYY" in zeilen
    assert "2026-10-18T12:00:00+00:00;standard;Fixterminüberwachung;H6;FALSCH;" in zeilen
    assert len(zeilen) == 1 + 6 + 1 + 1
    with pytest.raises(ValueError):
        excel.schreibe_bereich(REF, "Fixterminüberwachung", "D6:K6", [["zu", "kurz"]])
    with pytest.raises(ValueError):
        excel.schreibe_bereich(REF, "Fixterminüberwachung", "D6:E6", [["a", "b"]], zahlenformat=[["x"]])


def _xlsx_schreiben(pfad: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aktionen"
    ws.append(["Kunde", "Lieferant", "Von", "Bis", "Rabatt %"])
    ws.append(["Spar", "Cide", date(2026, 11, 1), date(2026, 11, 30), 0.1])
    ws["A4"] = "=SUMME(E2:E3)"  # Formel: mit data_only ohne gecachten Wert -> leer
    wb.create_sheet("Fixterminüberwachung").append(["KW"])
    wb.save(pfad)


def test_fake_excel_aus_xlsx_laedt_kopie_und_schreibt_nie_zurueck(tmp_path: Path):
    pfad = tmp_path / "Aktionszeiträume.xlsx"
    _xlsx_schreiben(pfad)
    vorher = pfad.read_bytes()
    blaetter = lade_xlsx(pfad)
    assert list(blaetter) == ["Aktionen", "Fixterminüberwachung"]
    assert blaetter["Aktionen"][1] == ["Spar", "Cide", datetime(2026, 11, 1), datetime(2026, 11, 30), 0.1]
    assert blaetter["Aktionen"][3][0] == ""  # Formel ohne Cache -> leer, nicht None
    excel = FakeExcelBackend(aus_xlsx=pfad)
    ref = DateiRefKonfig(lokaler_pfad=pfad, blatt="Aktionen")
    assert excel.blattnamen(ref) == ["Aktionen", "Fixterminüberwachung"]
    assert excel.blattnamen(DateiRefKonfig()) == ["Aktionen", "Fixterminüberwachung"]  # auch als "standard"
    excel.schreibe_bereich(ref, "Aktionen", "A3:B3", [["Rheinberg", "Bollina"]])
    assert excel.lese_bereich(ref, "Aktionen", "A3:B3") == [["Rheinberg", "Bollina"]]
    assert pfad.read_bytes() == vorher  # Datei unverändert
    # lazy: Referenz mit lokaler_pfad wird beim ersten Zugriff geladen
    excel2 = FakeExcelBackend()
    assert excel2.lese_bereich(ref, "Aktionen", "A2") == [["Spar"]]
    with pytest.raises(M365NichtGefunden):
        lade_xlsx(tmp_path / "fehlt.xlsx")


# ---------------------------------------------------------------- Mail


def entwurf(*anhaenge: Path, **kw) -> EmailEntwurf:
    felder = dict(
        typ=EntwurfTyp.SPEDITION_AUFTRAG,
        an=["auftraege@example.org"],
        cc=["kollegin@example.org"],
        betreff="Neuer Auftrag Beispiel – Abholung Cide / KW 44?",
        text_html="<html><body><p>Guten Tag Frau Beispiel,</p><p>anbei der Speditionsauftrag &amp; Lieferschein.</p></body></html>",
        text_plain="Guten Tag Frau Beispiel,\n\nanbei der Speditionsauftrag & Lieferschein.",
        anhaenge=list(anhaenge),
        absender="logistik@example.org",
    )
    felder.update(kw)
    return EmailEntwurf(**felder)


@pytest.fixture()
def mail_backend(tmp_path: Path) -> FakeMailBackend:
    return FakeMailBackend(tmp_path / "entwuerfe", jetzt=lambda: datetime(2026, 10, 18, 14, 5, 9).astimezone())


def test_fake_mail_entwurf_als_eml(mail_backend: FakeMailBackend, tmp_path: Path):
    pdf = tmp_path / "Spedition Beispiel – Speditionsauftrag – Cide – Beispiel – 18.10.2026.pdf"
    pdf.write_bytes(b"%PDF-1.4 test")
    e = entwurf(pdf, tmp_path / "fehlt.pdf")
    pfad = Path(mail_backend.erstelle_entwurf(e))
    assert pfad.parent == tmp_path / "entwuerfe"
    assert pfad.name == "20261018_140509_spedition_auftrag_Neuer Auftrag Beispiel – Abholung Cide _ KW 44_.eml"
    assert e.speicherort == str(pfad) and e.erstellt_am is not None and mail_backend.entwuerfe == [e]
    assert e.hinweise == [f"Anhang fehlt und wurde nicht beigefügt: {tmp_path / 'fehlt.pdf'}"]
    roh = pfad.read_bytes()
    assert b"\r\n" in roh
    msg = email.message_from_bytes(roh, policy=email.policy.default)
    assert msg["X-Unsent"] == "1" and msg["Subject"] == e.betreff
    assert msg["Message-ID"].startswith("<") and msg["Message-ID"].endswith("@example.org>")
    assert msg["From"] == "logistik@example.org" and msg["To"] == "auftraege@example.org" and msg["Cc"] == "kollegin@example.org"
    assert msg["X-Auftragsautomat-Typ"] == "spedition_auftrag" and msg["Date"]
    assert msg.get_content_type() == "multipart/mixed"
    assert msg.get_body(preferencelist=("plain",)).get_content().replace("\r\n", "\n").strip() == e.text_plain
    assert msg.get_body(preferencelist=("html",)).get_content().replace("\r\n", "\n").strip() == e.text_html
    anhaenge = list(msg.iter_attachments())
    assert len(anhaenge) == 1 and anhaenge[0].get_filename() == pdf.name and anhaenge[0].get_content() == b"%PDF-1.4 test"
    assert anhaenge[0].get_content_type() == "application/pdf"
    # zweiter Entwurf mit gleichem Namen bekommt " (2)"
    pfad2 = Path(mail_backend.erstelle_entwurf(entwurf()))
    assert pfad2.name.endswith(" (2).eml") and pfad2.exists()


def test_fake_mail_entwurf_ohne_absender_und_plain(mail_backend: FakeMailBackend):
    e = entwurf(absender=None, cc=[], text_plain="")
    msg = email.message_from_bytes(Path(mail_backend.erstelle_entwurf(e)).read_bytes(), policy=email.policy.default)
    assert msg["From"] is None and msg["Cc"] is None
    assert msg["Message-ID"].endswith("@auftragsautomat.invalid>")
    assert "Guten Tag Frau Beispiel," in msg.get_body(preferencelist=("plain",)).get_content()
    assert msg.get_content_type() == "multipart/alternative"


def _posteingang(tmp_path: Path) -> list[Mail]:
    def mail(id_: str, betreff: str, absender: str, empfangen: datetime, ordner: str = "inbox", anhaenge: list[MailAnhang] | None = None) -> Mail:
        return Mail(
            kopf=MailKopf(id=id_, betreff=betreff, absender=absender, empfangen=empfangen, hat_anhaenge=bool(anhaenge), ordner=ordner),
            text=f"Text zu {betreff}",
            anhaenge=anhaenge or [],
        )

    (tmp_path / "fixtures").mkdir(parents=True)
    (tmp_path / "fixtures" / "Bestätigung.pdf").write_bytes(b"pdf")
    (tmp_path / "direkt.pdf").write_bytes(b"direkt")
    return [
        mail("M1", "Auftragsbestätigung; Your Ref: B-4711; Our Ref: DEMB00000", "orders@example.org", datetime(2026, 10, 18, 9, 0, tzinfo=UTC),
             anhaenge=[MailAnhang(id="A1", name="Bestätigung.pdf"), MailAnhang(id=str(tmp_path / "direkt.pdf"), name="direkt.pdf")]),
        mail("M2", "Order confirmation", "wein@example.org", datetime(2026, 10, 17, 9, 0, tzinfo=UTC)),
        mail("M3", "Alt", "orders@example.org", datetime(2026, 9, 1, 9, 0, tzinfo=UTC)),
        mail("M4", "Our Ref: DEMB00001", "orders@example.org", datetime(2026, 10, 18, 10, 0, tzinfo=UTC), ordner="Spedition"),
    ]


def test_fake_mail_posteingang(tmp_path: Path):
    backend = FakeMailBackend(tmp_path / "e", _posteingang(tmp_path), anhaenge_ordner=tmp_path / "fixtures")
    seit = datetime(2026, 10, 11, tzinfo=UTC)
    assert [k.id for k in backend.suche_mails(seit=seit)] == ["M4", "M1", "M2"]
    assert [k.id for k in backend.suche_mails(seit=seit, absender="ORDERS@example.org")] == ["M4", "M1"]
    assert [k.id for k in backend.suche_mails(seit=seit, betreff_enthaelt="our ref")] == ["M4", "M1"]
    assert [k.id for k in backend.suche_mails(seit=seit, ordner="inbox")] == ["M1", "M2"]
    assert [k.id for k in backend.suche_mails(seit=seit, maximal=1)] == ["M4"]
    mail = backend.lese_mail("M1")
    assert mail.text == "Text zu Auftragsbestätigung; Your Ref: B-4711; Our Ref: DEMB00000"
    mail.anhaenge.clear()
    assert len(backend.lese_mail("M1").anhaenge) == 2  # Kopie, Original unverändert
    with pytest.raises(M365NichtGefunden):
        backend.lese_mail("M9")
    (tmp_path / "ziel").mkdir()
    ziel = backend.speichere_anhang("M1", "A1", tmp_path / "ziel")  # vorhandener Ordner -> Anhangsname
    assert ziel == tmp_path / "ziel" / "Bestätigung.pdf" and ziel.read_bytes() == b"pdf"
    assert backend.speichere_anhang("M1", "Bestätigung.pdf", tmp_path / "ziel" / "neu.pdf").read_bytes() == b"pdf"
    assert backend.speichere_anhang("M1", "direkt.pdf", tmp_path / "ziel").read_bytes() == b"direkt"
    with pytest.raises(M365NichtGefunden):
        backend.speichere_anhang("M1", "A9", tmp_path / "ziel")
    ohne_ordner = FakeMailBackend(tmp_path / "e", _posteingang(tmp_path / "zwei"))
    with pytest.raises(M365NichtGefunden, match="Datei für Anhang"):
        ohne_ordner.speichere_anhang("M1", "A1", tmp_path / "ziel")


def test_mail_aus_eml(mail_backend: FakeMailBackend, tmp_path: Path, fixtures_ordner: Path):
    pdf = tmp_path / "Beleg.pdf"
    pdf.write_bytes(b"%PDF")
    pfad = Path(mail_backend.erstelle_entwurf(entwurf(pdf, absender="Spedition Beispiel <info@example.org>")))
    mail = mail_aus_eml(pfad, mail_id="X1", ordner="Spedition")
    assert mail.kopf.id == "X1" and mail.kopf.absender == "info@example.org" and mail.kopf.absender_name == "Spedition Beispiel"
    assert mail.kopf.betreff.startswith("Neuer Auftrag Beispiel") and mail.kopf.ordner == "Spedition" and mail.kopf.hat_anhaenge
    assert mail.kopf.empfangen.tzinfo is not None
    assert mail.an == ["auftraege@example.org"] and mail.cc == ["kollegin@example.org"]
    assert mail.text.startswith("Guten Tag Frau Beispiel") and "<p>" in mail.html
    assert [(a.id, a.name, a.content_type) for a in mail.anhaenge] == [("Beleg.pdf", "Beleg.pdf", "application/pdf")]
    # vorhandene Fixture-Mail des Projekts (Bestellung an das Weingut) ist ebenfalls lesbar
    fixture = fixtures_ordner / "emails" / "bestellung_cide_mit_pdf_anhang.eml"
    if fixture.exists():
        f = mail_aus_eml(fixture)
        assert f.kopf.id and f.kopf.betreff and f.text


# ---------------------------------------------------------------- Dateien


def test_fake_datei_backend(tmp_path: Path):
    backend = FakeDateiBackend(tmp_path / "dossiers")
    ref = DateiRefKonfig()
    assert not backend.existiert(ref, "Beispiel – WIK – B-4711")
    ordner = backend.ordner_anlegen(ref, "Beispiel – WIK – B-4711/06_Mails")
    assert Path(ordner).is_dir() and backend.existiert(ref, "Beispiel – WIK – B-4711/06_Mails") and backend.existiert(ref)
    assert backend.ordner_anlegen(ref, "Beispiel – WIK – B-4711/06_Mails") == ordner  # idempotent
    quelle = tmp_path / "Cde.pdf"
    quelle.write_bytes(b"%PDF")
    hochgeladen = backend.hochladen(quelle, ref, "Beispiel – WIK – B-4711/")
    assert Path(hochgeladen) == tmp_path / "dossiers" / "Beispiel – WIK – B-4711" / "Cde.pdf"
    assert Path(backend.hochladen(quelle, ref, "Beispiel – WIK – B-4711/02_Bestellung.pdf")).name == "02_Bestellung.pdf"
    assert backend.hochladen(quelle, ref, "") == str(tmp_path / "dossiers" / "Cde.pdf")
    text = backend.schreibe_text(ref, "Beispiel – WIK – B-4711/07_Fixtermin.txt", "KW 43 | Cide | Beispiel ä")
    assert Path(text).read_text(encoding="utf-8") == "KW 43 | Cide | Beispiel ä"
    eintraege = backend.liste(ref, "Beispiel – WIK – B-4711")
    assert [(e.name, e.pfad, e.ist_ordner) for e in eintraege] == [
        ("02_Bestellung.pdf", "Beispiel – WIK – B-4711/02_Bestellung.pdf", False),
        ("06_Mails", "Beispiel – WIK – B-4711/06_Mails", True),
        ("07_Fixtermin.txt", "Beispiel – WIK – B-4711/07_Fixtermin.txt", False),
        ("Cde.pdf", "Beispiel – WIK – B-4711/Cde.pdf", False),
    ]
    assert eintraege[0].groesse == 4 and eintraege[1].groesse is None and eintraege[0].geaendert.tzinfo is not None
    assert [e.pfad for e in backend.liste(ref)] == ["Beispiel – WIK – B-4711", "Cde.pdf"]
    (tmp_path / "runter").mkdir()
    ziel = backend.herunterladen(ref, "Beispiel – WIK – B-4711/Cde.pdf", tmp_path / "runter")
    assert ziel == tmp_path / "runter" / "Cde.pdf" and ziel.read_bytes() == b"%PDF"
    assert backend.herunterladen(ref, "Cde.pdf", tmp_path / "runter2" / "x.pdf").read_bytes() == b"%PDF"
    with pytest.raises(M365NichtGefunden):
        backend.herunterladen(ref, "fehlt.pdf", tmp_path)
    with pytest.raises(M365NichtGefunden):
        backend.liste(ref, "fehlt")
    with pytest.raises(M365Fehler, match="verlässt"):
        backend.existiert(ref, "../../etc")
    with pytest.raises(M365Fehler, match="als Datei"):
        backend.ordner_anlegen(ref, "Cde.pdf")
    with pytest.raises(M365Fehler, match="Lokale Datei fehlt"):
        backend.hochladen(tmp_path / "nix.pdf", ref, "")
    # Referenz mit eigenem lokalen Pfad gewinnt über die Wurzel
    eigene = DateiRefKonfig(lokaler_pfad=tmp_path / "sync" / "Aufträge")
    assert backend.ordner_anlegen(eigene, "Aufträge ohne Provision") == str(tmp_path / "sync" / "Aufträge" / "Aufträge ohne Provision")


# ---------------------------------------------------------------- Fabrik


def test_fabrik_fake(settings: Settings):
    backends = fabrik.erzeuge_backends(settings)
    assert backends.backend == "fake" and backends.anmeldung is None
    assert isinstance(backends.excel, FakeExcelBackend) and isinstance(backends.mail, FakeMailBackend)
    assert isinstance(backends.datei, FakeDateiBackend)
    assert backends.mail.entwuerfe_ordner == settings.kommunikation.entwuerfe_ordner
    assert backends.datei.wurzel == settings.dossier.lokaler_ordner
    assert fabrik.pruefe_anmeldung(settings) == "Backend 'fake': keine Microsoft-365-Anmeldung nötig."
    assert fabrik.anmelden(settings).startswith("Backend 'fake'")


def test_fabrik_lokal(settings: Settings):
    settings.m365.backend = "lokal"
    backends = fabrik.erzeuge_backends(settings)
    from auftragsautomat.m365.lokal import ExcelComBackend, LokalerDateiBackend, OutlookComMailBackend

    assert backends.backend == "lokal"
    assert isinstance(backends.excel, ExcelComBackend) and isinstance(backends.mail, OutlookComMailBackend)
    assert isinstance(backends.datei, LokalerDateiBackend) and backends.datei.wurzel == settings.dossier.lokaler_ordner


def test_fabrik_graph_ohne_konfiguration(settings: Settings, tmp_path: Path):
    settings.m365.backend = "graph"
    assert fabrik.pruefe_anmeldung(settings).startswith("Graph nicht konfiguriert: settings.m365.client_id fehlt")
    with pytest.raises(M365Fehler, match="client_id"):
        fabrik.erzeuge_backends(settings)


def test_fabrik_graph_nicht_angemeldet(settings: Settings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    import msal

    class FakeApp:
        def __init__(self, *a, **kw) -> None:
            pass

        def get_accounts(self, username=None):
            return []

    monkeypatch.setattr(msal, "PublicClientApplication", FakeApp)
    settings.m365.backend = "graph"
    settings.m365.client_id = "00000000-0000-0000-0000-0000000000aa"
    settings.m365.tenant_id = "00000000-0000-0000-0000-000000000001"
    settings.m365.token_cache = tmp_path / "cache.bin"
    meldung = fabrik.pruefe_anmeldung(settings)
    assert meldung.startswith("Nicht angemeldet") and "m365-anmelden" in meldung
    backends = fabrik.erzeuge_backends(settings)
    from auftragsautomat.m365.graph import GraphDateiBackend, GraphExcelBackend, GraphMailBackend

    assert backends.backend == "graph" and backends.anmeldung is not None
    assert isinstance(backends.excel, GraphExcelBackend) and isinstance(backends.mail, GraphMailBackend)
    assert isinstance(backends.datei, GraphDateiBackend) and backends.mail.postfach is None


def test_fabrik_graph_angemeldet(settings: Settings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    import msal

    class FakeApp:
        def __init__(self, *a, **kw) -> None:
            self.ausgaben: list[str] = []

        def get_accounts(self, username=None):
            return [{"username": "frau.beispiel@example.org"}]

        def acquire_token_silent(self, scopes, account=None, **kw):
            return {"access_token": "tok"}

        def initiate_device_flow(self, scopes=None, **kw):
            return {"user_code": "ABCD", "message": "Code ABCD eingeben"}

        def acquire_token_by_device_flow(self, flow, **kw):
            return {"access_token": "tok", "id_token_claims": {"preferred_username": "frau.beispiel@example.org"}}

    monkeypatch.setattr(msal, "PublicClientApplication", FakeApp)
    settings.m365.backend = "graph"
    settings.m365.client_id = "00000000-0000-0000-0000-0000000000aa"
    settings.m365.tenant_id = "00000000-0000-0000-0000-000000000001"
    settings.m365.token_cache = tmp_path / "cache.bin"
    assert fabrik.pruefe_anmeldung(settings) == (
        f"Angemeldet als frau.beispiel@example.org (Tenant {settings.m365.tenant_id}, Cache {tmp_path / 'cache.bin'})."
    )
    ausgaben: list[str] = []
    assert fabrik.anmelden(settings, device_code=True, ausgabe=ausgaben.append).startswith("Angemeldet als frau.beispiel@example.org")
    assert ausgaben == ["Code ABCD eingeben"]


def test_fabrik_unbekanntes_backend(settings: Settings):
    settings.m365.backend = "irgendwas"  # type: ignore[assignment]
    with pytest.raises(M365Fehler, match="Unbekanntes"):
        fabrik.erzeuge_backends(settings)
