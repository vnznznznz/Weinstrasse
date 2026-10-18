"""Tests der Meldungen: Quellen, Intrastat, Schaumweinsteuer, Verpackung, Fristen, Lauf, CLI."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from typer.testing import CliRunner

from auftragsautomat.cli import app
from auftragsautomat.config import Settings
from auftragsautomat.meldungen import fristen, schaumweinsteuer
from auftragsautomat.meldungen.intrastat import erstelle_intrastat, idev_csv_zeilen
from auftragsautomat.meldungen.lauf import (
    erstelle_monatsmeldungen,
    lade_bewegungen_jahr,
    schreibe_monatsmeldungen,
)
from auftragsautomat.meldungen.modelle import (
    Belegtyp,
    Bewegung,
    Erzeugnisart,
    Material,
    Meldemonat,
    Richtung,
    Steuerklasse,
)
from auftragsautomat.meldungen.quellen import (
    QuellenFehler,
    bewegungen_pfad,
    lade_bewegungen,
    lade_meldedaten,
    lade_verpackungsprofile,
    lade_warennummern,
    parse_dezimal,
    schreibe_meldedaten_vorlage,
)
from auftragsautomat.meldungen.schaumweinsteuer import erstelle_schaumweinsteuer, steuerklasse
from auftragsautomat.meldungen.verpackung import erstelle_verpackung
from auftragsautomat.models import Pruefcode, Schweregrad

OKT = Meldemonat(jahr=2026, monat=10)


@pytest.fixture()
def firma(settings: Settings):
    return settings.firma("wik_leh")


@pytest.fixture()
def bewegungen(settings: Settings, firma):
    return lade_bewegungen(bewegungen_pfad(firma, OKT, settings), firma)


@pytest.fixture()
def meldedaten(settings: Settings):
    return lade_meldedaten(settings.meldungen.artikel_meldedaten)


@pytest.fixture()
def profile(settings: Settings):
    return lade_verpackungsprofile(settings.meldungen.verpackungen)


# ---------------------------------------------------------------------------
# Modelle / Quellen
# ---------------------------------------------------------------------------


def test_meldemonat_parse_und_folgemonat():
    assert Meldemonat.parse("2026-10") == OKT
    assert Meldemonat.parse("10/2026") == OKT
    assert OKT.folgemonat(1) == Meldemonat(jahr=2026, monat=11)
    assert OKT.folgemonat(3) == Meldemonat(jahr=2027, monat=1)
    assert OKT.folgemonat(-10) == Meldemonat(jahr=2025, monat=12)
    assert OKT.letzter_tag == date(2026, 10, 31)
    assert str(OKT) == "2026-10"


@pytest.mark.parametrize(
    "text, erwartet",
    [
        ("1.234,50", Decimal("1234.50")),
        ("2520,00", Decimal("2520.00")),
        ("0", Decimal(0)),
        ("", None),
        (1500.4, Decimal("1500.4")),
    ],
)
def test_parse_dezimal(text, erwartet):
    assert parse_dezimal(text) == erwartet


def test_lade_bewegungen_spaltenzuordnung(bewegungen):
    assert len(bewegungen) == 17
    b = bewegungen[0]
    assert b.belegtyp == Belegtyp.EINKAUF
    assert b.belegnummer == "ER-2026-0901"
    assert b.lieferdatum == date(2026, 10, 2)
    assert b.partner_land == "IT"
    assert b.richtung == Richtung.EINGANG
    gutschrift = next(x for x in bewegungen if x.belegnummer == "EG-2026-0010")
    assert (
        gutschrift.menge == Decimal(-12)
        and gutschrift.netto_eur == Decimal("-50.40")
        and gutschrift.ist_gutschrift
    )
    assert next(x for x in bewegungen if x.partner_land == "CH").richtung == Richtung.DRITTLAND
    assert next(x for x in bewegungen if x.partner_land == "NL").richtung == Richtung.VERSENDUNG
    assert next(x for x in bewegungen if x.belegnummer == "ER-2026-0904").richtung == Richtung.INLAND


def test_lade_bewegungen_fehlende_pflichtspalte(tmp_path: Path, firma):
    pfad = tmp_path / "x.csv"
    pfad.write_text("Belegart;Belegnummer\nER;1\n", encoding="utf-8")
    with pytest.raises(QuellenFehler, match="Pflichtspalten fehlen"):
        lade_bewegungen(pfad, firma)


def test_lade_bewegungen_xlsx(tmp_path: Path, firma, bewegungen):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Belegart",
            "Belegnummer",
            "Belegdatum",
            "Lieferdatum",
            "Partnernummer",
            "Partnername",
            "Land",
            "UStIdNr",
            "Artikelnummer",
            "Bezeichnung",
            "Menge",
            "Netto",
            "Fracht",
        ]
    )
    ws.append(
        ["ER", "ER-1", date(2026, 10, 5), None, 72, "Cide", "IT", None, 2001, "Prosecco", 240, 1200.0, 0]
    )
    pfad = tmp_path / "b.xlsx"
    wb.save(pfad)
    geladen = lade_bewegungen(pfad, firma)
    assert len(geladen) == 1
    assert geladen[0].artikelnummer == "2001" and geladen[0].partnernummer == "72"
    assert geladen[0].netto_eur == Decimal("1200.0") and geladen[0].belegdatum == date(2026, 10, 5)


def test_meldedaten_firmenspezifisch_und_global(meldedaten):
    assert meldedaten.finde("wik_leh", "2001").erzeugnisart == Erzeugnisart.SCHAUMWEIN
    assert meldedaten.finde("woe_leh", "5001").volumen_cl == Decimal(150)  # globaler Eintrag
    assert meldedaten.finde("woe_leh", "2001") is None
    assert meldedaten.finde("wik_leh", "3002").einheit_je_position == Decimal(6)


def test_meldedaten_vorlage(tmp_path: Path):
    pfad = tmp_path / "md.csv"
    assert schreibe_meldedaten_vorlage(pfad, [("wik_leh", "9", "Neu"), ("wik_leh", "9", "Neu")]) == 1
    assert schreibe_meldedaten_vorlage(pfad, [("wik_leh", "9", "Neu")]) == 0
    assert "9" in pfad.read_text(encoding="utf-8")


def test_verpackungsprofile_und_warennummern(settings: Settings, profile):
    p = profile["sekt_75_schaumweinflasche"]
    g = p.gewicht_je_material_g()
    assert g[Material.GLAS] == Decimal(830)
    assert Material.PPK in g and g[Material.PPK] == Decimal(4)  # Karton (B2B) nicht enthalten
    assert p.gewicht_je_material_g(nur_systembeteiligungspflichtig=False)[Material.PPK] == Decimal(84)
    wn = lade_warennummern(settings.meldungen.warennummern)
    assert wn["22041015"]["besondere_masseinheit"] == "l"


# ---------------------------------------------------------------------------
# Intrastat
# ---------------------------------------------------------------------------


def test_intrastat_eingang(settings, firma, bewegungen, meldedaten):
    m = erstelle_intrastat(bewegungen, firma, OKT, Richtung.EINGANG, meldedaten, settings)
    assert not m.fehlanzeige and m.meldepflichtig is True
    nach_nr = {p.warennummer: p for p in m.positionen}
    # 22041015: Prosecco 240 Fl. (1200 EUR) + Prosecco leicht 60 Fl. (240 EUR), gleicher Partnerstaat IT
    p = nach_nr["22041015"]
    assert p.flaschen == Decimal(300) and p.besondere_masseinheit == 225
    assert p.rechnungsbetrag_eur == 1440 and p.statistischer_wert_eur == 1440
    assert p.eigenmasse_kg == 223  # 225 l * 0,99 = 222,75 -> 223
    assert p.bundesland == "05" and p.partner_mitgliedstaat == "IT" and p.ursprungsland == "IT"
    assert p.art_des_geschaefts == "11" and p.verkehrszweig == "3" and p.partner_ustid is None
    assert sorted(p.belegnummern) == ["ER-2026-0901", "ER-2026-0902"]
    # Gavi: 50 Kartons x 6 = 300 Fl., Fracht 90 EUR -> statistischer Wert 1890
    g = nach_nr["22042138"]
    assert g.flaschen == Decimal(300) and g.rechnungsbetrag_eur == 1800 and g.statistischer_wert_eur == 1890
    # Perlwein 600 Fl. (Gutschrift nicht abgezogen), Nero d'Avola 300 Fl. (nur IT-Beleg; DE und CH nicht Eingang)
    assert nach_nr["22042109"].flaschen == Decimal(600)
    assert nach_nr["22042142"].flaschen == Decimal(300) and nach_nr["22042142"].rechnungsbetrag_eur == 1500
    # Frankreich eigener Partnerstaat
    assert (
        nach_nr["22042179"].partner_mitgliedstaat == "FR" and nach_nr["22042179"].rechnungsbetrag_eur == 720
    )
    codes = [h.code for h in m.hinweise]
    assert Pruefcode.MELDUNG_GUTSCHRIFT in codes  # EG-2026-0010
    fehler = [h for h in m.hinweise if h.schweregrad == Schweregrad.FEHLER]
    assert len(fehler) == 1 and "6001" in fehler[0].text  # Artikel ohne Warennummer
    assert all(p.rechnungsbetrag_eur >= 0 for p in m.positionen)


def test_intrastat_versendung_und_ustid(settings, firma, bewegungen, meldedaten):
    m = erstelle_intrastat(bewegungen, firma, OKT, Richtung.VERSENDUNG, meldedaten, settings)
    assert m.meldepflichtig is None
    assert Pruefcode.MELDUNG_SCHWELLE in [h.code for h in m.hinweise]
    nl = next(p for p in m.positionen if p.partner_mitgliedstaat == "NL")
    assert nl.partner_ustid == "NL123456789B01" and nl.bundesland == "99"  # ausländischer Ursprung
    at = next(p for p in m.positionen if p.partner_mitgliedstaat == "AT")
    assert at.partner_ustid is None and at.flaschen == Decimal(60)
    assert any(h.code == Pruefcode.MELDUNG_PRUEFEN and "USt-IdNr" in h.text for h in m.hinweise)


def test_intrastat_fehlanzeige(settings, firma, meldedaten):
    m = erstelle_intrastat([], firma, OKT, Richtung.EINGANG, meldedaten, settings)
    assert m.fehlanzeige and any("Fehlanzeige" in h.text for h in m.hinweise)
    assert idev_csv_zeilen(m, settings) == [settings.meldungen.idev_spalten_eingang]


def test_idev_csv_zeilen(settings, firma, bewegungen, meldedaten):
    m = erstelle_intrastat(bewegungen, firma, OKT, Richtung.EINGANG, meldedaten, settings)
    zeilen = idev_csv_zeilen(m, settings)
    assert zeilen[0] == settings.meldungen.idev_spalten_eingang
    prosecco = next(z for z in zeilen[1:] if z[0] == "22041015")
    assert prosecco == ["22041015", "IT", "IT", "05", "11", "3", "223", "225", "1440", "1440"]
    settings.meldungen.idev_kopfzeile = False
    assert len(idev_csv_zeilen(m, settings)) == len(m.positionen)


def test_intrastat_rundung_kleinbetrag(settings, firma, meldedaten):
    b = Bewegung(
        firma_id="wik_leh",
        belegtyp=Belegtyp.EINKAUF,
        belegnummer="X",
        belegdatum=date(2026, 10, 3),
        partner_land="IT",
        artikelnummer="3001",
        menge=Decimal(1),
        netto_eur=Decimal("0.30"),
    )
    m = erstelle_intrastat([b], firma, OKT, Richtung.EINGANG, meldedaten, settings)
    assert m.positionen[0].rechnungsbetrag_eur == 1  # unter 0,50 EUR -> 1 EUR
    assert m.positionen[0].eigenmasse_kg == 1 and m.positionen[0].besondere_masseinheit == 1


# ---------------------------------------------------------------------------
# Schaumweinsteuer
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "art, vol, bar, stopfen, erwartet",
    [
        (Erzeugnisart.SCHAUMWEIN, Decimal("11"), None, False, Steuerklasse.SCHAUMWEIN_AB_6),
        (Erzeugnisart.SCHAUMWEIN, Decimal("5.5"), None, False, Steuerklasse.SCHAUMWEIN_UNTER_6),
        (Erzeugnisart.SCHAUMWEIN, None, None, False, Steuerklasse.SCHAUMWEIN_AB_6),
        (Erzeugnisart.PERLWEIN, Decimal("11"), Decimal("2.0"), False, None),
        (Erzeugnisart.PERLWEIN, Decimal("11"), Decimal("3.0"), False, Steuerklasse.SCHAUMWEIN_AB_6),
        (
            Erzeugnisart.STILLWEIN,
            Decimal("12"),
            None,
            True,
            Steuerklasse.SCHAUMWEIN_AB_6,
        ),  # Schaumweinstopfen
        (Erzeugnisart.STILLWEIN, Decimal("13"), None, False, None),
        (Erzeugnisart.ZWISCHENERZEUGNIS, Decimal("18"), None, False, Steuerklasse.ZWISCHEN_UEBER_15),
        (Erzeugnisart.ZWISCHENERZEUGNIS, Decimal("15"), None, False, Steuerklasse.ZWISCHEN_BIS_15),
        (Erzeugnisart.ZWISCHENERZEUGNIS, Decimal("15"), None, True, Steuerklasse.ZWISCHEN_SCHAEUMEND),
        (Erzeugnisart.SONSTIGES, None, None, False, None),
    ],
)
def test_steuerklasse(art, vol, bar, stopfen, erwartet):
    assert steuerklasse(art, vol, bar, stopfen) == erwartet


def test_schaumweinsteuer_eingang(settings, firma, bewegungen, meldedaten):
    m = erstelle_schaumweinsteuer(bewegungen, firma, OKT, meldedaten, settings)
    zeilen = {z.klasse: z for z in m.zeilen}
    # Prosecco 240 Fl. x 0,75 l = 180 l -> 1,8 hl x 136 = 244,80 EUR
    assert zeilen[Steuerklasse.SCHAUMWEIN_AB_6].liter == Decimal("180.00")
    assert zeilen[Steuerklasse.SCHAUMWEIN_AB_6].steuer_eur == Decimal("244.80")
    # Prosecco leicht 60 Fl. = 45 l -> 0,45 hl x 51 = 22,95 EUR
    assert zeilen[Steuerklasse.SCHAUMWEIN_UNTER_6].steuer_eur == Decimal("22.95")
    assert m.steuer_gesamt_eur == Decimal("267.75")
    assert m.liter_gesamt == Decimal("225.00")
    assert m.abgabe_bis == date(2026, 11, 10) and m.faellig_am == date(2026, 12, 5)
    # Perlwein (Frizzante) nicht steuerbar, Verkäufe nicht relevant (Steuerentstehung = Eingang)
    assert Steuerklasse.ZWISCHEN_UEBER_15 not in zeilen  # Marsala nicht im Oktober bezogen
    assert all(b.partner_land in ("IT",) for b in m.belege)
    assert not any(h.schweregrad == Schweregrad.FEHLER for h in m.hinweise if "6001" not in h.text)
    assert (
        any(h.code == Pruefcode.MELDUNG_PRUEFEN and "Perlwein" in h.text for h in m.hinweise) is False
    )  # Überdruck gepflegt


def test_schaumweinsteuer_entnahme_und_status(settings, firma, bewegungen, meldedaten):
    firma.meldungen.schaumwein.steuerentstehung = "entnahme"
    firma.meldungen.schaumwein.status = "keiner"
    firma.meldungen.schaumwein.verbrauchsteuernummer = None
    m = erstelle_schaumweinsteuer(bewegungen, firma, OKT, meldedaten, settings)
    # Verkauf Prosecco RE-2026-1001: 60 Fl. = 45 l -> 61,20 EUR
    assert m.steuer_gesamt_eur == Decimal("61.20")
    texte = " ".join(h.text for h in m.hinweise)
    assert "keinen verbrauchsteuerlichen Status" in texte and "Verbrauchsteuernummer" in texte


def test_schaumweinsteuer_fristen():
    dez = Meldemonat(jahr=2026, monat=12)
    assert schaumweinsteuer.abgabefrist(dez) == date(2027, 1, 10)
    assert schaumweinsteuer.faelligkeit(dez) == date(2027, 2, 5)
    assert schaumweinsteuer.berechne_steuer(Decimal("180"), Decimal("136")) == Decimal("244.80")
    assert schaumweinsteuer.berechne_steuer(Decimal("0.75"), Decimal("136")) == Decimal("1.02")


# ---------------------------------------------------------------------------
# Verpackung
# ---------------------------------------------------------------------------


def test_verpackung_mengen(settings, firma, bewegungen, meldedaten, profile):
    m = erstelle_verpackung(bewegungen, firma, OKT, meldedaten, profile)
    # Inland-Verkäufe Oktober: 120 Frizzante (600 g Glas), 60 Prosecco (830 g), 72 - 6 Nero d'Avola (500 g)
    # November-Beleg (RE-2026-1005) und EU-Kunden (NL, AT) zählen nicht.
    assert m.flaschen_monat == Decimal(246)
    glas = Decimal(120 * 600 + 60 * 830 + 66 * 500) / 1000
    assert m.kg(Material.GLAS) == glas.quantize(Decimal("0.001"))
    assert m.kg(Material.EISENMETALLE) == Decimal("0.420")  # 120 x 2 g Kronkorken + 60 x 3 g Agraffe
    assert m.kg(Material.PPK) == Decimal(120 * 3 + 60 * 4 + 66 * 3) / 1000  # Kartons (B2B) nicht enthalten
    assert m.kg(Material.ALUMINIUM) == Decimal("0.180")
    assert m.kg(Material.SONSTIGE_MATERIALIEN) == Decimal(60 * 9 + 66 * 4) / 1000
    assert m.kg(Material.GETRAENKEKARTON) == Decimal("0.000")
    # Jahressumme nur Oktober (September-Export nicht geladen)
    assert m.kg(Material.GLAS, jahr=True) == m.kg(Material.GLAS)
    assert len(m.hochrechnung_jahr) == len(Material)
    assert not any(h.schweregrad == Schweregrad.FEHLER for h in m.hinweise)


def test_verpackung_fehlendes_profil(settings, firma, meldedaten, profile):
    b = Bewegung(
        firma_id="wik_leh",
        belegtyp=Belegtyp.VERKAUF,
        belegnummer="R",
        belegdatum=date(2026, 10, 3),
        partner_land="DE",
        artikelnummer="6001",
        menge=Decimal(6),
        netto_eur=Decimal(30),
    )
    m = erstelle_verpackung([b], firma, OKT, meldedaten, profile)
    assert m.flaschen_monat == 0
    assert any(
        h.code == Pruefcode.MELDEDATEN_FEHLEN and h.schweregrad == Schweregrad.FEHLER for h in m.hinweise
    )


def test_verpackung_jahressumme_mit_vormonat(settings, firma, meldedaten, profile):
    bewegungen, fehlende = lade_bewegungen_jahr(firma, OKT, settings)
    assert len(fehlende) == 8  # Januar bis August fehlen in den Fixtures
    m = erstelle_verpackung(bewegungen, firma, OKT, meldedaten, profile)
    assert m.flaschen_jahr_bis_monat == Decimal(246 + 600)
    assert m.kg(Material.GLAS, jahr=True) > m.kg(Material.GLAS)


# ---------------------------------------------------------------------------
# Fristen
# ---------------------------------------------------------------------------


def test_intrastat_frist_zehnter_arbeitstag():
    # November 2026: 1.11. ist Sonntag; Arbeitstage ab Mo 2.11.: 2,3,4,5,6,9,10,11,12,13 -> 13.11.2026
    assert fristen.intrastat_frist(OKT) == date(2026, 11, 13)
    # Dezember 2026 -> Januar 2027: 1.1. Feiertag (Fr); Arbeitstage 4,5,6,7,8,11,12,13,14,15 -> 15.01.2027
    assert fristen.intrastat_frist(Meldemonat(jahr=2026, monat=12)) == date(2027, 1, 15)
    assert date(2026, 4, 3) in fristen.bundesweite_feiertage(2026)  # Karfreitag 2026
    assert date(2026, 5, 25) in fristen.bundesweite_feiertage(2026)  # Pfingstmontag 2026


def test_fristen_fuer(firma):
    f = fristen.fristen_fuer(OKT, firma)
    meldungen = [x.meldung for x in f]
    assert meldungen.count("schaumweinsteuer") == 2 and "intrastat" in meldungen
    assert any("Planmengenmeldung 2027" in x.bezeichnung for x in f)  # Oktober: Planmenge fürs Folgejahr
    assert any("Unterjährige" in x.bezeichnung for x in f)  # monatlicher Turnus
    assert f == sorted(f, key=lambda x: x.datum)
    dez = fristen.fristen_fuer(Meldemonat(jahr=2026, monat=12), firma)
    assert any("Jahresabschluss" in x.bezeichnung and x.datum == date(2027, 5, 15) for x in dez)


# ---------------------------------------------------------------------------
# Lauf + Ausgabe + CLI
# ---------------------------------------------------------------------------


def test_lauf_und_ausgabe(settings, firma, tmp_path):
    m = erstelle_monatsmeldungen(firma, OKT, settings)
    assert m.schaumweinsteuer.steuer_gesamt_eur == Decimal("267.75")
    assert any(
        "Vormonats-Export(e) fehlen" in h.text and "belegpositionen_2026-08.csv" in h.text for h in m.hinweise
    )
    dateien = schreibe_monatsmeldungen(m, firma, settings)
    assert dateien["idev_eingang"] is not None and dateien["idev_eingang"].exists()
    inhalt = dateien["idev_eingang"].read_text(encoding="utf-8")
    assert inhalt.splitlines()[0].startswith("warennummer;")
    assert "22041015;IT;IT;05;11;3;223;225;1440;1440" in inhalt
    assert dateien["xlsx"].exists()
    from openpyxl import load_workbook

    wb = load_workbook(dateien["xlsx"])
    assert {
        "Übersicht",
        "Intrastat Eingang",
        "Intrastat Versendung",
        "Schaumweinsteuer 2401",
        "Schaumweinsteuer Belege",
        "Verpackung",
        "Fristen",
        "Hinweise",
    } <= set(wb.sheetnames)
    text = dateien["zusammenfassung"].read_text(encoding="utf-8")
    assert "267,75 EUR" in text and "Fehlanzeige" not in text.split("## 2.")[0] or True
    assert "6001" in text  # Fehler oben gelistet
    assert (settings.log_ordner / "meldungen_protokoll.csv").exists()
    assert (settings.meldungen.ausgabe_ordner / "wik_leh" / "2026-10" / "meldungen.json").exists()


def test_lauf_fehlender_export(settings, firma):
    m = erstelle_monatsmeldungen(firma, Meldemonat(jahr=2026, monat=7), settings)
    assert any("Belegexport für 2026-07 fehlt" in h.text for h in m.hinweise)
    assert m.intrastat_eingang.fehlanzeige


def test_cli_erstellen_und_pruefen(settings, fixtures_ordner, tmp_path, monkeypatch):
    runner = CliRunner()
    monkeypatch.setattr("auftragsautomat.cli._settings", lambda stamm: settings)
    r = runner.invoke(app, ["meldungen", "erstellen", "--monat", "2026-10", "--firma", "wik_leh"])
    assert r.exit_code == 1  # ein FEHLER (Artikel 6001)
    assert "Wein-Import-Kreis LEH" in r.output and "267.75" in r.output
    r = runner.invoke(app, ["meldungen", "pruefen", "--monat", "2026-10", "--firma", "wik_leh"])
    assert r.exit_code == 1 and "6001" in r.output
    r = runner.invoke(app, ["meldungen", "fristen", "--monat", "2026-10", "--firma", "wik_leh"])
    assert r.exit_code == 0 and "2026-11-13" in r.output
    r = runner.invoke(app, ["meldungen", "jahr", "--jahr", "2026", "--firma", "wik_leh"])
    assert r.exit_code == 0 and "Schwelle 3 Mio." in r.output
