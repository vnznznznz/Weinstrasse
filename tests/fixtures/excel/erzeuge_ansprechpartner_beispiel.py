"""Erzeugt die synthetische Kontaktliste ``ansprechpartner_beispiel.xlsx`` für die Tests.

Aufbau wie die echte "Ansprechpartner Abwicklung.xlsx" (docs/AUSBAUSTUFE_2.md, Abschnitt 1 und 2.3):

* Blatt "Übersicht" (wird ignoriert),
* Blatt "Lieferanten": Zeile 1 Titel, dann Blöcke je Land (Landname in Spalte A, darunter Kopf
  ``Firma | Name | Position | E-Mail | Telefon``, Zeilen, Leerzeile trennt),
* Blatt "Speditionen": Titel, Erklärung, dann je Spedition eine Tabelle (Tabellenobjekt
  ``tbl_Spedition_<Name>``) mit Kopf ``Name | Bereich / Position | E-Mail | Telefon`` – der
  Speditionsname steht nicht als Zelle über der Tabelle,
* Blatt "Kunden": Titel, dann Blöcke je Kunde (Kundenname in Spalte A, Kopf
  ``Produkt | Name | Position | E-Mail | Telefon | Spedition | Abwicklungsverlauf``).

Alle Namen, Adressen und Nummern sind erfunden (Domain example.org). Eingebaute Stolpersteine:
Telefonnummer als Zahl, Mehrfach-Leerzeichen im Namen, fehlende Spaltenwerte, leere Firma
(gilt wie die Zeile darüber), Mehrfachfirma mit "&", Spedition "<Name> (für <Kunde>)".

Aufruf: ``python tests/fixtures/excel/erzeuge_ansprechpartner_beispiel.py [zielpfad]``
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

ZIEL = Path(__file__).with_name("ansprechpartner_beispiel.xlsx")

KOPF_LIEFERANTEN = ["Firma", "Name", "Position", "E-Mail", "Telefon"]
KOPF_SPEDITIONEN = ["Name", "Bereich / Position", "E-Mail", "Telefon"]
KOPF_KUNDEN = ["Produkt", "Name", "Position", "E-Mail", "Telefon", "Spedition", "Abwicklungsverlauf"]

LIEFERANTEN: list[list[Any]] = [
    ["Ansprechpartner Lieferanten"],
    [],
    ["Italien"],
    KOPF_LIEFERANTEN,
    ["Weingut Beispiel Srl", "Giulia   Esempio", "Orders", "orders@weingut-beispiel.example.org", "+39 0123 456789"],
    ["", "Marco Esempio", "Export Manager", "export@weingut-beispiel.example.org", 390123456780],
    [
        "Cantina Alpha & Cantina Beta & Cantina Gamma",
        "Anna Esempio",
        "Orders / Logistica",
        "ordini@cantine-abc.example.org",
        "+39 0456 112233",
    ],
    [],
    ["Deutschland"],
    KOPF_LIEFERANTEN,
    ["Weingut Musterhof GmbH", "Max Muster", "Vertrieb", "vertrieb@musterhof.example.org", "0049 6131 1234567"],
    ["Weingut Musterhof GmbH", "Erika Muster", "Buchhaltung", "buchhaltung@musterhof.example.org", None],
    [],
    ["Frankreich"],
    KOPF_LIEFERANTEN,
    ["Domaine Exemple", "Jean Exemple", "Orders", "commandes@domaine-exemple.example.org", "+33 4 12 34 56 78"],
    ["Domaine Exemple", "Marie Exemple", None, "marie@domaine-exemple.example.org", None],
]

SPEDITIONEN: list[list[Any]] = [
    ["Ansprechpartner Speditionen"],
    ["Neue Aufträge ausschließlich an das Auftrags-Postfach senden; Statusfragen an das Info-Postfach."],
    [],
    KOPF_SPEDITIONEN,  # Zeile 4: tbl_Spedition_Hillebrand (A4:D8)
    ["Auftrags-Postfach", "Postfach – Neue Aufträge (ausschließlich hierhin)", "neue-auftraege@hillebrand.example.org", None],
    ["Info-Postfach", "Sendungsnachfragen, Status", "info@hillebrand.example.org", None],
    ["Petra Beispiel", "Disponentin", "p.beispiel@hillebrand.example.org", "+49 6131 9876543"],
    ["Karl Beispiel", "EMCS / Zoll", "zoll@hillebrand.example.org", "+49 6131 9876544"],
    [],
    KOPF_SPEDITIONEN,  # Zeile 10: tbl_Spedition_Beispiel_Logistik (A10:D12)
    ["Zentrale", "Disposition", "dispo@beispiel-logistik.example.org", "0049 201 123456"],
    ["Rechnungen", "Buchhaltung", "buchhaltung@beispiel-logistik.example.org", None],
]

KUNDEN: list[list[Any]] = [
    ["Ansprechpartner Kunden"],
    [],
    ["Beispielmarkt Nord"],
    KOPF_KUNDEN,
    [
        "Allgemein",
        "Frau Beispiel",
        "Disposition",
        "dispo@beispielmarkt.example.org",
        "+49 201 555 0100",
        "Beispiel-Logistik (für Beispielmarkt Nord)",
        "Bestellung per Mail; Auftragskopie an die Dispo",
    ],
    ["Weingut Beispiel", "Herr Beispiel", "Einkauf Wein Italien", "einkauf@beispielmarkt.example.org", None, None, None],
    ["Allgemein", "Buchhaltung", "Buchhaltung", "rechnung@beispielmarkt.example.org", None, None, None],
    [],
    ["Weinhandlung Beispiel"],
    KOPF_KUNDEN,
    ["Allgemein", "Hans Beispiel", "Inhaber", "info@weinhandlung-beispiel.example.org", "0201 5550200", None, "Abholung durch Spedition"],
]


def _schreibe(ws: Any, zeilen: list[list[Any]]) -> None:
    for zeile in zeilen:
        ws.append(zeile if zeile else [None])


def erzeuge(ziel: Path = ZIEL) -> Path:
    """Schreibt die Beispiel-Kontaktliste nach ``ziel`` und liefert den Pfad."""
    mappe = Workbook()
    uebersicht = mappe.active
    uebersicht.title = "Übersicht"
    uebersicht.append(["Übersicht"])
    uebersicht.append(["Diese Mappe enthält die Ansprechpartner für Lieferanten, Speditionen und Kunden."])

    lieferanten = mappe.create_sheet("Lieferanten")
    _schreibe(lieferanten, LIEFERANTEN)

    speditionen = mappe.create_sheet("Speditionen")
    _schreibe(speditionen, SPEDITIONEN)
    stil = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    for name, ref in (("tbl_Spedition_Hillebrand", "A4:D8"), ("tbl_Spedition_Beispiel_Logistik", "A10:D12")):
        tabelle = Table(displayName=name, ref=ref)
        tabelle.tableStyleInfo = stil
        speditionen.add_table(tabelle)

    kunden = mappe.create_sheet("Kunden")
    _schreibe(kunden, KUNDEN)

    ziel = Path(ziel)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    mappe.save(ziel)
    return ziel


if __name__ == "__main__":
    pfad = erzeuge(Path(sys.argv[1]) if len(sys.argv) > 1 else ZIEL)
    print(f"Kontaktliste geschrieben: {pfad}")
