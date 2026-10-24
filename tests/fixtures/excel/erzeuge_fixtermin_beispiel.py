"""Erzeugt ``fixterminueberwachung_beispiel.xlsx`` – ein synthetisches Abbild der Fixterminüberwachung.

Aufruf (aus dem Projektstamm)::

    python tests/fixtures/excel/erzeuge_fixtermin_beispiel.py

Struktur wie die echte Datei (docs/AUSBAUSTUFE_2.md, Abschnitt 1):

* Blatt "Fixterminüberwachung": Zeile 1 Titel, Zeile 2 Legende, Zeile 3 leer, Zeile 4 Kopfzeile
  (KW | Datum | Tag | Lieferant | Kunde | Auftragsdatum | Verladedatum | Bestätigung des
  Lieferanten | Bestätigung der Spedition | Abgeholt | Bemerkung), ab Zeile 5 je Werktag fünf
  Zeilen (Slots). In den ersten vier Slots verweisen die Spalten A–C per Formel auf die
  Ankerzeile des Tages (der fünfte Slot); hier stehen die Formeln **als Text** (``=A9``), damit
  openpyxl sie beim Lesen nicht als unausgewertete Formel (leer) liefert. Die Ankerzeile trägt
  KW, Datum (echtes Datum) und Tag. Feiertage sind eine einzelne Zeile mit Bemerkung
  "Feiertag: …". Spalten H–J sind boolesche Werte (in Excel 365 Kontrollkästchen).
* Blatt "Ressourcen für Auto DropDown": die Kurznamen-Listen (in der echten Datei per
  LET/SORT/UNIQUE/FILTER-Formel gefüllt, hier als Werte; nie beschreiben).

Zeitraum: KW 20–23/2026 (11.05.–05.06.2026) mit Christi Himmelfahrt (14.05.), Pfingstmontag
(25.05.) und Fronleichnam (04.06., nur in einigen Bundesländern, u. a. NW). Alle Namen, Nummern und
Texte sind synthetisch.
"""

from __future__ import annotations

import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HIER = Path(__file__).resolve().parent
ZIEL = HIER / "fixterminueberwachung_beispiel.xlsx"
PROJEKT = HIER.parents[2]
if str(PROJEKT / "src") not in sys.path:
    sys.path.insert(0, str(PROJEKT / "src"))

from auftragsautomat.termine.kalender import feiertag, ist_wochenende, kw, wochentag_kurz  # noqa: E402

BLATT = "Fixterminüberwachung"
RESSOURCEN = "Ressourcen für Auto DropDown"
KOPF = [
    "KW",
    "Datum",
    "Tag",
    "Lieferant",
    "Kunde",
    "Auftragsdatum",
    "Verladedatum",
    "Bestätigung des Lieferanten",
    "Bestätigung der Spedition",
    "Abgeholt",
    "Bemerkung",
]
KOPFZEILE = 4
SLOTS = 5
BUNDESLAND = "NW"
VON = date(2026, 5, 11)  # Montag KW 20
BIS = date(2026, 6, 5)  # Freitag KW 23

# Vorbelegte Slots (synthetisch): Datum -> Liste je Slot
# (Lieferant, Kunde, Auftragsdatum, Verladedatum, Lieferant best., Spedition best., Abgeholt, Bemerkung)
Eintrag = tuple[str, str, date | None, date | None, bool, bool, bool, str]
EINTRAEGE: dict[date, list[Eintrag]] = {
    date(2026, 5, 12): [
        ("Beispielgut", "Musterhandel", date(2026, 5, 4), date(2026, 5, 11), True, False, False, "Best.-Nr 20260001, AB 30260001, Spedition Beispiel"),
        ("Musterweine", "Beispielmarkt", date(2026, 5, 5), None, False, False, False, ""),
    ],
    date(2026, 5, 13): [
        ("Musterweine", "Musterhandel", date(2026, 4, 28), date(2026, 5, 8), True, True, True, "Best.-Nr 20260000, AB 30260000, Spedition Beispiel | Ref. DEMB00000"),
    ],
    date(2026, 5, 20): [
        ("Beispielgut", "Markt Eins", date(2026, 5, 6), None, False, False, False, ""),
        ("Beispielgut", "Markt Zwei", date(2026, 5, 6), None, False, False, False, ""),
        ("Beispielgut", "Markt Drei", date(2026, 5, 6), None, False, False, False, ""),
        ("Musterweine", "Markt Vier", date(2026, 5, 7), None, False, False, False, ""),
        ("Musterweine", "Markt Fünf", date(2026, 5, 7), None, False, False, False, ""),
    ],
    date(2026, 6, 2): [
        ("Beispielgut", "Musterhandel", date(2026, 5, 18), None, False, False, False, "Best.-Nr 20260002, AB 30260002, Spedition Beispiel"),
    ],
}
LIEFERANTEN = ["Beispielgut", "Musterweine", "Weingut Beispiel"]
KUNDEN = ["Musterhandel", "Beispielmarkt", "Markt Eins", "Markt Zwei", "Markt Drei", "Markt Vier", "Markt Fünf"]


def _text_formel(zelle, formel: str) -> None:
    """Formel als Text ablegen (openpyxl würde '=A9' sonst als Formel speichern)."""
    zelle.value = formel
    zelle.data_type = "s"


def erzeuge(ziel: Path = ZIEL) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = BLATT
    fett = Font(bold=True)
    ws["A1"] = "Fixterminüberwachung – synthetisches Beispiel (Testdaten)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Legende: je Werktag fünf Zeilen; Spalten A–C per Formel aus der letzten Zeile des Tages; "
        "H–J Kontrollkästchen (WAHR/FALSCH); Feiertage als eigene Zeile mit Bemerkung."
    )
    for spalte, titel in enumerate(KOPF, start=1):
        zelle = ws.cell(row=KOPFZEILE, column=spalte, value=titel)
        zelle.font = fett
        zelle.fill = PatternFill("solid", fgColor="DDEBF7")
        zelle.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[KOPFZEILE].height = 32

    zeile = KOPFZEILE + 1
    tag = VON
    while tag <= BIS:
        if ist_wochenende(tag):
            tag += timedelta(days=1)
            continue
        name = feiertag(tag, BUNDESLAND)
        if name:
            ws.cell(row=zeile, column=1, value=kw(tag))
            ws.cell(row=zeile, column=2, value=tag).number_format = "DD.MM.YYYY"
            ws.cell(row=zeile, column=3, value=wochentag_kurz(tag))
            ws.cell(row=zeile, column=11, value=f"Feiertag: {name}")
            for spalte in range(1, 12):
                ws.cell(row=zeile, column=spalte).fill = PatternFill("solid", fgColor="EDEDED")
            zeile += 1
            tag += timedelta(days=1)
            continue
        anker = zeile + SLOTS - 1
        eintraege = EINTRAEGE.get(tag, [])
        for slot in range(SLOTS):
            z = zeile + slot
            if slot < SLOTS - 1:
                _text_formel(ws.cell(row=z, column=1), f"=A{anker}")
                _text_formel(ws.cell(row=z, column=2), f"=B{anker}")
                _text_formel(ws.cell(row=z, column=3), f"=C{anker}")
            else:
                ws.cell(row=z, column=1, value=kw(tag))
                ws.cell(row=z, column=2, value=tag).number_format = "DD.MM.YYYY"
                ws.cell(row=z, column=3, value=wochentag_kurz(tag))
            if slot < len(eintraege):
                lieferant, kunde, auftrag, verladung, best_l, best_s, abgeholt, bemerkung = eintraege[slot]
                ws.cell(row=z, column=4, value=lieferant)
                ws.cell(row=z, column=5, value=kunde)
                if auftrag is not None:
                    ws.cell(row=z, column=6, value=auftrag).number_format = "DD.MM.YYYY"
                if verladung is not None:
                    ws.cell(row=z, column=7, value=verladung).number_format = "DD.MM.YYYY"
                ws.cell(row=z, column=8, value=best_l)
                ws.cell(row=z, column=9, value=best_s)
                ws.cell(row=z, column=10, value=abgeholt)
                if bemerkung:
                    ws.cell(row=z, column=11, value=bemerkung)
            else:
                ws.cell(row=z, column=8, value=False)
                ws.cell(row=z, column=9, value=False)
                ws.cell(row=z, column=10, value=False)
        zeile += SLOTS
        tag += timedelta(days=1)

    breiten = [6, 12, 5, 16, 16, 14, 14, 14, 14, 10, 60]
    for i, breite in enumerate(breiten, start=1):
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = ws.cell(row=KOPFZEILE + 1, column=1)

    ws2 = wb.create_sheet(RESSOURCEN)
    ws2["A1"], ws2["B1"] = "Lieferanten", "Kunden"
    ws2["A1"].font = ws2["B1"].font = fett
    for i, name in enumerate(LIEFERANTEN, start=2):
        ws2.cell(row=i, column=1, value=name)
    for i, name in enumerate(KUNDEN, start=2):
        ws2.cell(row=i, column=2, value=name)
    ws2.column_dimensions["A"].width = ws2.column_dimensions["B"].width = 20

    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)
    return ziel


if __name__ == "__main__":
    pfad = erzeuge()
    print(f"geschrieben: {pfad}")
